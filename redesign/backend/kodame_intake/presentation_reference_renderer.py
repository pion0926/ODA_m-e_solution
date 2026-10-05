"""Editable report-style layouts reconstructed from the reviewed PDF samples.

PDFs do not contain PowerPoint masters. All report text and tables are native
PPTX objects; the two samples' aspect ratios and hierarchy are explicit here.
"""
import math
import re
import unicodedata
from io import BytesIO

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Pt
from pptx.oxml.xmlchemy import OxmlElement
from .presentation_profiles import table_widths


def lines(text, width, size):
    return sum(max(1, math.ceil(sum(1 if unicodedata.east_asian_width(c) in "WF" else .55 for c in line) * size / width))
               for line in str(text).split("\n"))


def text(slide, x, y, w, h, value, size=15, bold=False, color="111111", align=PP_ALIGN.LEFT):
    if lines(value, w-4, size) * size * 1.25 > h:
        raise ValueError(f"텍스트 상자 높이 초과: {str(value)[:45]} (문장을 압축해 주세요)")
    shape = slide.shapes.add_textbox(Pt(x), Pt(y), Pt(w), Pt(h))
    frame = shape.text_frame
    frame.word_wrap = True
    frame.margin_left = frame.margin_right = Pt(2)
    frame.margin_top = frame.margin_bottom = Pt(0)
    p = frame.paragraphs[0]
    p.text = str(value)
    p.font.name = "나눔고딕"
    p.font.size = Pt(size)
    p.font.bold = bold
    p.font.color.rgb = RGBColor.from_string(color)
    p.line_spacing = 1.15
    p.space_after = Pt(0)
    p.alignment = align
    return shape


def rectangle(slide, x, y, w, h, color):
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Pt(x), Pt(y), Pt(w), Pt(h))
    shape.fill.solid()
    shape.fill.fore_color.rgb = RGBColor.from_string(color)
    shape.line.fill.background()
    # Suppress the template theme's default drop shadow on flat rules/bands.
    shape._element.spPr.append(OxmlElement('a:effectLst'))
    for effect in shape._element.xpath('./p:style/a:effectRef'):
        effect.set('idx', '0')
    return shape


def picture(slide, photo, x, y, w, h):
    scale = min(w / photo["width"], h / photo["height"])
    pw, ph = photo["width"]*scale, photo["height"]*scale
    slide.shapes.add_picture(BytesIO(photo["data"]), Pt(x+(w-pw)/2), Pt(y+(h-ph)/2), Pt(pw), Pt(ph))


def table(slide, columns, rows, w, brief, top=104):
    count = len(columns)
    widths = table_widths(columns, w)
    size = 16 if count <= 3 else 15
    heights = [max(38, max(lines(c, widths[j]-14, size) for j, c in enumerate(row))*size*1.25+16) for row in rows]
    if sum(heights) > 356:
        raise ValueError("표 본문이 356pt를 초과합니다. 근거를 유지하며 각 셀을 더 짧게 작성해 주세요.")
    obj = slide.shapes.add_table(len(rows)+1, count, Pt(28), Pt(top), Pt(w), Pt(34+sum(heights))).table
    for j, width in enumerate(widths):
        obj.columns[j].width = Pt(width)
    for i, row in enumerate([columns]+rows):
        obj.rows[i].height = Pt(34 if i == 0 else heights[i-1])
        for j, value in enumerate(row):
            cell = obj.cell(i, j)
            cell.text = str(value)
            cell.margin_left = cell.margin_right = Pt(7)
            cell.margin_top = cell.margin_bottom = Pt(5)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            cell.fill.solid()
            cell.fill.fore_color.rgb = RGBColor.from_string(("24466A" if brief else "DCEAF3") if i == 0 else ("F1F5F8" if i % 2 else "FFFFFF"))
            for p in cell.text_frame.paragraphs:
                p.font.name = "나눔고딕"
                p.font.size = Pt(size)
                p.font.bold = i == 0 or j == 0
                p.font.color.rgb = RGBColor.from_string("FFFFFF" if i == 0 and brief else "111111")
                p.line_spacing = 1.1
                p.space_after = Pt(0)


def build_reference_deck(profile, slides, source, photos, partial=False):
    if not partial and len(slides) != profile["slide_count"]:
        raise ValueError("선택한 장수와 생성된 장수가 일치하지 않습니다.")
    prs = Presentation()
    width = profile["width_pt"]
    prs.slide_width, prs.slide_height = Pt(width), Pt(540)
    brief = profile["slide_count"] == 15
    for page, item in zip(profile["pages"], slides):
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        layout = page["layout"]
        if layout == "cover":
            if brief:
                rectangle(slide, 0, 0, width, 540, "2B5688")
                rectangle(slide, 72, 215, width-144, 62, "FFFFFF")
                text(slide, 78, 225, width-156, 45, page["title"], 28, True, "243E61", PP_ALIGN.CENTER)
                text(slide, 60, 120, width-120, 85, source["summary"]["project"]["title"], 19, True, "FFFFFF", PP_ALIGN.CENTER)
                text(slide, 80, 315, width-160, 95, str(source["summary"]["project"].get("period") or ""), 16, False, "FFFFFF", PP_ALIGN.CENTER)
            else:
                text(slide, 36, 60, width-72, 60, page["title"], 32, True)
                text(slide, 36, 135, width-72, 105, source["summary"]["project"]["title"], 23, True)
                if item.get("photo_id"):
                    picture(slide, photos[item["photo_id"]], 36, 255, width-72, 195)
                rectangle(slide, 28, 487, width-56, 2, "3F4143")
        else:
            repeated_heading = not page['section'] or re.sub(r'^[ⅠⅡⅢⅣⅤⅥ.\s]+', '', page['section']) == page['title']
            if brief:
                rectangle(slide, 18, 18, width-36, 37, "315E91")
                text(slide, 30, 25, width-60, 28, page["section"] or page["title"], 19, True, "FFFFFF")
                if not repeated_heading:
                    text(slide, 28, 68, width-56, 30, page["title"], 19, True, "243E61")
            else:
                text(slide, 26, 19, width*.65, 29, page["title"], 19, True)
                text(slide, width*.68, 25, width*.28, 22, page["section"], 11, True, "163E61", PP_ALIGN.RIGHT)
                rectangle(slide, 26, 53, width-52, 3, "414141")
                rectangle(slide, width*.22, 53, width*.24, 3, "00739D")
            if layout == "toc":
                entries = []
                for p in profile["pages"]:
                    if p["section"] and p["section"] not in [e[0] for e in entries]:
                        entries.append((p["section"], p["slide_number"]))
                for i, (name, number) in enumerate(entries):
                    y = 128+i*54
                    text(slide, 90, y, width-210, 32, name, 19, True, "274D79" if brief else "64006B")
                    text(slide, width-105, y, 40, 32, str(number), 18, align=PP_ALIGN.RIGHT)
            elif page["columns"]:
                table(slide, page["columns"], item["rows"], width-56, brief, top=82 if brief and repeated_heading else 104)
            else:
                blocks = item.get("blocks", [])
                has_photo = bool(item.get("photo_id")) and layout == "photo"
                body_w = (width-76)/2 if has_photo else width-64
                for i, block in enumerate(blocks):
                    y = 109+i*(180 if has_photo else 128)
                    text(slide, 30, y, body_w, 30, block["heading"], 17, True, "176598")
                    text(slide, 32, y+34, body_w-4, 131 if has_photo else 90, block["text"], 16)
                if has_photo:
                    picture(slide, photos[item["photo_id"]], width/2+12, 116, body_w-10, 287)
                    text(slide, width/2+12, 421, body_w-10, 63, item.get("caption", ""), 12, color="454545")
            rectangle(slide, 26, 513, width-52, 1, "AAAAAA")
            text(slide, width/2-40, 518, 80, 20, f"{page['slide_number']} / {profile['slide_count']}", 10, align=PP_ALIGN.CENTER)
        note = item.get("speaker_notes", "") + "\n[Sources]\n" + "\n".join(item["source_sections"])
        if item.get("photo_id"):
            photo = photos[item["photo_id"]]
            note += f"\n사진: {photo['file_name']}, p. {photo['page']}"
        slide.notes_slide.notes_text_frame.text = note
    out = BytesIO()
    prs.save(out)
    return out.getvalue()
