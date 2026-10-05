from __future__ import annotations
from .llm_models import current_llm_model
from .model_catalog import prepare_model_payload

import json
import re
import uuid
from io import BytesIO
from pathlib import Path
from typing import Any

import httpx
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_AUTO_SHAPE_TYPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt
from psycopg.types.json import Jsonb

from .db import connection, tenant_context
from .presentation_prompt import (
    ALLOWED_LAYOUTS,
    LAYOUT_SEQUENCE,
    PLAN_SCHEMA,
    ROLE_SEQUENCE,
    ROLE_SOURCE_FALLBACKS,
    ROLE_TEXT_BUDGETS,
    SLIDE_COUNT,
    SYSTEM_PROMPT,
    presentation_prompt,
)
from .presentation_quality import (
    render_and_validate_presentation,
    score_presentation_quality,
    validate_presentation_bytes,
)
from .presentation_reference import build_reference_context
from .presentation_source import collect_presentation_source
from .report_text import sanitize_report_text
from .settings import (
    OPENROUTER_API_KEY,
    OPENROUTER_BASE_URL,
    OPENROUTER_REFERER,
    PRESENTATION_MAX_GENERATION_ATTEMPTS,
    PRESENTATION_MIN_QUALITY_SCORE,
)
from .usage import record_token_usage
from .ai.request_limits import request_slot


PRESENTATION_DIR = Path("/app/data/presentation_exports")

PALETTES = {
    "ocean": {
        "ink": "17324D", "navy": "0C3B5D", "blue": "2E7DA7", "teal": "25A6A1",
        "orange": "E8864A", "sand": "F2E8D7", "paper": "F7F8F6", "white": "FFFFFF",
        "muted": "667A89", "line": "D6E1E6", "pale": "E7F3F5",
    },
    "forest": {
        "ink": "203833", "navy": "153F3A", "blue": "4A778D", "teal": "2F8C72",
        "orange": "D58A4B", "sand": "EEE2CF", "paper": "F7F7F2", "white": "FFFFFF",
        "muted": "687A72", "line": "D8E1DC", "pale": "E7F1EC",
    },
    "ink": {
        "ink": "20283A", "navy": "18213A", "blue": "526FAF", "teal": "3A9E9B",
        "orange": "E4774E", "sand": "EFE4D9", "paper": "F7F6F4", "white": "FFFFFF",
        "muted": "6C7381", "line": "D9DDE5", "pale": "E9EDF5",
    },
    "cobalt": {
        "ink": "172A46", "navy": "0B3564", "blue": "1F68A7", "teal": "2B98A4",
        "orange": "DF7546", "sand": "F0E7DA", "paper": "F6F8FA", "white": "FFFFFF",
        "muted": "60758A", "line": "D5E0EA", "pale": "E6F0F7",
    },
}


def _presentation_prompt(
    source: dict[str, Any],
    reference_context: dict[str, Any] | None = None,
    qa_feedback: list[str] | None = None,
) -> str:
    return presentation_prompt(source, reference_context, qa_feedback)


def _clean_text(value: Any, limit: int) -> str:
    text = sanitize_report_text(str(value or ""))
    text = re.sub(r"\s+", " ", text).strip()
    if len(text) > limit:
        text = text[: limit - 1].rstrip(" ,.;:") + "…"
    return text


def _clean_list(value: Any, *, limit: int = 5, item_limit: int = 150) -> list[str]:
    if not isinstance(value, list):
        return []
    result = [_clean_text(item, item_limit) for item in value]
    return [item for item in result if item][:limit]


def _clean_metrics(value: Any) -> list[dict[str, str]]:
    if not isinstance(value, list):
        return []
    result: list[dict[str, str]] = []
    for item in value[:4]:
        if not isinstance(item, dict):
            continue
        metric = {
            "value": _clean_text(item.get("value"), 18),
            "label": _clean_text(item.get("label"), 55),
            "context": _clean_text(item.get("context"), 78),
        }
        if metric["value"] or metric["label"]:
            result.append(metric)
    return result


def _roadmap_step_parts(value: Any) -> tuple[str, list[str]]:
    """Turn Claude's compact roadmap contract into a fit-safe card."""
    parts = [part.strip() for part in re.split(r"\s*\|\s*", _clean_text(value, 130)) if part.strip()]
    title = _clean_text(parts[0] if parts else value, 26) or "후속 과제"
    if len(parts) < 2:
        return title, ["담당·시점·확인지표 지정"]
    labels = ("담당", "시점", "확인")
    details = []
    for index, label in enumerate(labels, 1):
        if index >= len(parts):
            break
        limit = 27 if label != "확인" else 34
        details.append(f"{label} · {_clean_text(parts[index], limit)}")
    if len(parts) > 4:
        details[-1] = _clean_text(f"{details[-1]} · {' · '.join(parts[4:])}", 42)
    return title, details


def _ordered_pdm_steps(item: dict[str, Any]) -> list[str]:
    """Guarantee the four PDM stages even when the model omits a stage."""
    raw_steps = [_clean_text(value, 70) for value in (item.get("steps") or []) if _clean_text(value, 70)]
    stage_specs = (
        ("투입·활동", ("투입", "활동"), "예산·인력·시설과 핵심 실행과제"),
        ("산출물", ("산출",), "사업이 직접 제공한 제품·서비스"),
        ("성과", ("성과",), "대상과 제도에서 확인된 변화"),
        ("영향", ("영향",), "장기 파급효과와 후속 검증 과제"),
    )
    ordered: list[str] = []
    used: set[int] = set()
    for label, keywords, fallback in stage_specs:
        match_index = next(
            (index for index, value in enumerate(raw_steps) if index not in used and any(key in value for key in keywords)),
            None,
        )
        if match_index is None:
            ordered.append(f"{label}: {fallback}")
            continue
        used.add(match_index)
        value = raw_steps[match_index]
        ordered.append(value if value.startswith(label) or (label == "산출물" and value.startswith("산출")) else f"{label}: {value}")
    return ordered


def _extract_json(text: str) -> dict[str, Any]:
    value = str(text or "").strip()
    if value.startswith("```"):
        value = re.sub(r"^```(?:json)?\s*|\s*```$", "", value, flags=re.I)
    if not value.startswith("{"):
        start, end = value.find("{"), value.rfind("}")
        if start >= 0 and end > start:
            value = value[start : end + 1]
    data = json.loads(value)
    if not isinstance(data, dict):
        raise ValueError("Claude 발표자료 응답이 JSON 객체가 아닙니다.")
    return data


def normalize_presentation_plan(
    raw: dict[str, Any],
    project: dict[str, Any],
    source: dict[str, Any] | None = None,
    reference_context: dict[str, Any] | None = None,
) -> dict[str, Any]:
    slides = raw.get("slides")
    if not isinstance(slides, list) or len(slides) != SLIDE_COUNT:
        raise ValueError(f"Claude 발표자료 구성은 정확히 {SLIDE_COUNT}장이어야 합니다.")

    allowed_sections = {
        str(item.get("part_id") or "") for item in (source or {}).get("report_sections", [])
    }
    allowed_documents = {
        str(item.get("file_name") or "") for item in (source or {}).get("evidence_catalog", [])
    }
    normalized: list[dict[str, Any]] = []
    for index, item in enumerate(slides):
        if not isinstance(item, dict):
            raise ValueError(f"{index + 1}번 슬라이드 구성이 객체가 아닙니다.")
        layout = str(item.get("layout") or "").strip()
        if layout not in ALLOWED_LAYOUTS:
            layout = LAYOUT_SEQUENCE[index]
        source_sections = _clean_list(item.get("source_sections"), limit=8, item_limit=70)
        source_documents = _clean_list(item.get("source_documents"), limit=12, item_limit=180)
        if allowed_sections:
            source_sections = [value for value in source_sections if value in allowed_sections]
        if allowed_documents:
            source_documents = [value for value in source_documents if value in allowed_documents]
        if allowed_sections and not source_sections:
            source_sections = [
                value for value in ROLE_SOURCE_FALLBACKS.get(ROLE_SEQUENCE[index], ())
                if value in allowed_sections
            ]
        if index > 0 and allowed_sections and not source_sections:
            raise ValueError(f"{index + 1}번 슬라이드의 유효한 보고서 근거 섹션이 없습니다.")
        role = ROLE_SEQUENCE[index]
        budget = ROLE_TEXT_BUDGETS[role]
        normalized.append({
            "slide_number": index + 1,
            "role": role,
            "layout": layout,
            "eyebrow": _clean_text(item.get("eyebrow"), 55),
            "title": _clean_text(item.get("title"), budget["title"]) or f"평가결과 {index + 1}",
            "headline": _clean_text(item.get("headline"), 75),
            "body": _clean_text(item.get("body"), 220),
            "bullets": _clean_list(
                item.get("bullets"),
                limit=budget["bullets"],
                item_limit=budget["bullet_chars"],
            ),
            "secondary_bullets": _clean_list(
                item.get("secondary_bullets"),
                limit=budget["bullets"],
                item_limit=budget["bullet_chars"],
            ),
            "metrics": _clean_metrics(item.get("metrics")),
            "metric_value": _clean_text(item.get("metric_value"), 28),
            "metric_label": _clean_text(item.get("metric_label"), 70),
            "quote": _clean_text(item.get("quote"), 150),
            "steps": _clean_list(item.get("steps"), limit=5, item_limit=70),
            "accent": str(item.get("accent") or "teal") if str(item.get("accent") or "") in {"blue", "teal", "orange"} else "teal",
            "visual_direction": _clean_text(item.get("visual_direction"), 200),
            "speaker_notes": _clean_text(item.get("speaker_notes"), 1200),
            "source_sections": source_sections,
            "source_documents": source_documents,
        })

    # Claude is asked to vary silhouettes, but the renderer also guarantees
    # diversity if a provider returns a repetitive but otherwise valid plan.
    if len({slide["layout"] for slide in normalized}) < 7 or any(
        normalized[index]["layout"] == normalized[index - 1]["layout"]
        for index in range(1, len(normalized))
    ):
        for index, slide in enumerate(normalized):
            slide["layout"] = LAYOUT_SEQUENCE[index]

    selected_variant = dict((reference_context or {}).get("selected_variant") or {})
    seed_text = str(selected_variant.get("seed") or "00")
    try:
        accent_offset = int(seed_text[:2], 16) % 3
    except ValueError:
        accent_offset = 0
    accent_cycle = ("blue", "teal", "orange")
    for index, slide in enumerate(normalized):
        slide["accent"] = accent_cycle[(index + accent_offset) % len(accent_cycle)]

    design = raw.get("design") if isinstance(raw.get("design"), dict) else {}
    palette = str(selected_variant.get("palette") or design.get("palette") or "ocean")
    if palette not in PALETTES:
        palette = "ocean"
    deck_title = _clean_text(raw.get("deck_title"), 90) or "종료평가 결과 브리핑"
    # Claude sometimes repeats the full project name as the deck title.  The
    # project name already appears as the subtitle; keeping the display title
    # short preserves a 50pt cover without collisions.
    if len(deck_title) > ROLE_TEXT_BUDGETS["cover"]["title"]:
        deck_title = "종료평가 결과 브리핑"
    return {
        "deck_title": deck_title,
        "subtitle": _clean_text(raw.get("subtitle"), 160) or _clean_text(project.get("title"), 160),
        "design": {
            "palette": palette,
            "mood": _clean_text(design.get("mood"), 100),
            "design_rationale": _clean_text(design.get("design_rationale"), 300),
            "variant_id": _clean_text(selected_variant.get("id"), 50) or "editorial_axis",
            "variant_seed": _clean_text(selected_variant.get("seed"), 20),
            "composition": _clean_text(selected_variant.get("composition"), 140),
            "title_treatment": _clean_text(selected_variant.get("title_treatment"), 120),
            "surface_treatment": _clean_text(selected_variant.get("surface_treatment"), 120),
            "reference_profile_version": _clean_text((reference_context or {}).get("profile_version"), 40),
        },
        "slides": normalized,
    }


def _collect_presentation_source() -> tuple[dict[str, Any], dict[str, Any], list[str]]:
    return collect_presentation_source()


def _message_content(response: httpx.Response) -> str:
    body = response.json()
    record_token_usage(body, current_llm_model())
    content = body["choices"][0]["message"]["content"]
    if isinstance(content, list):
        return "".join(str(item.get("text") or "") for item in content if isinstance(item, dict))
    return str(content or "")


def request_presentation_plan(
    source: dict[str, Any],
    reference_context: dict[str, Any] | None = None,
    qa_feedback: list[str] | None = None,
) -> dict[str, Any]:
    if not OPENROUTER_API_KEY:
        raise RuntimeError("OPENROUTER_API_KEY가 설정되지 않아 Claude 발표자료를 생성할 수 없습니다.")
    headers = {
        "Authorization": f"Bearer {OPENROUTER_API_KEY}",
        "Content-Type": "application/json",
        "HTTP-Referer": OPENROUTER_REFERER,
        "X-Title": "KODAME Executive Evaluation Presentation",
    }
    prompt = _presentation_prompt(source, reference_context, qa_feedback)
    formats = [
        {"type": "json_schema", "json_schema": PLAN_SCHEMA},
        {"type": "json_object"},
    ]
    last_error: Exception | None = None
    with httpx.Client(timeout=httpx.Timeout(300.0, connect=20.0)) as client:
        for response_format in formats:
            seed_text = str(((reference_context or {}).get("selected_variant") or {}).get("seed") or "0")
            try:
                generation_seed = int(seed_text[:8], 16)
            except ValueError:
                generation_seed = 0
            payload = {
                "model": current_llm_model(),
                "messages": [
                    {"role": "system", "content": SYSTEM_PROMPT},
                    {"role": "user", "content": prompt},
                ],
                "reasoning": {"effort": "high", "exclude": True},
                "temperature": 0.55,
                "seed": generation_seed,
                "max_completion_tokens": 14000,
                "response_format": response_format,
            }
            try:
                with request_slot(payload) as reservation:
                    response = client.post(f"{OPENROUTER_BASE_URL}/chat/completions", headers=headers, json=prepare_model_payload(payload))
                    from .ai.global_budget import settle
                    settle(reservation, response.json())
                if response.status_code >= 400:
                    raise RuntimeError(f"OpenRouter Claude 호출 실패: HTTP {response.status_code} {response.text[:300]}")
                return _extract_json(_message_content(response))
            except (RuntimeError, ValueError, KeyError, IndexError, TypeError, json.JSONDecodeError) as exc:
                last_error = exc
    raise RuntimeError(f"Claude가 유효한 {SLIDE_COUNT}장 발표자료 JSON을 반환하지 못했습니다.") from last_error


def _rgb(value: str) -> RGBColor:
    return RGBColor.from_string(value)


def _shape(slide, x: float, y: float, w: float, h: float, fill: str, *, radius: bool = False, line: str | None = None):
    kind = MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE if radius else MSO_AUTO_SHAPE_TYPE.RECTANGLE
    shape = slide.shapes.add_shape(kind, Inches(x), Inches(y), Inches(w), Inches(h))
    shape.fill.solid()
    shape.fill.fore_color.rgb = _rgb(fill)
    shape.line.color.rgb = _rgb(line or fill)
    if radius:
        shape.adjustments[0] = 0.12
    return shape


def _text(
    slide,
    x: float,
    y: float,
    w: float,
    h: float,
    value: str,
    *,
    size: float = 18,
    color: str = "17324D",
    bold: bool = False,
    align: PP_ALIGN = PP_ALIGN.LEFT,
    valign: MSO_ANCHOR = MSO_ANCHOR.TOP,
    margin: float = 0,
):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = box.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.vertical_anchor = valign
    frame.margin_left = frame.margin_right = frame.margin_top = frame.margin_bottom = Inches(margin)
    paragraph = frame.paragraphs[0]
    paragraph.text = value or ""
    paragraph.alignment = align
    paragraph.font.name = "맑은 고딕"
    paragraph.font.size = Pt(size)
    paragraph.font.bold = bold
    paragraph.font.color.rgb = _rgb(color)
    return box


def _bullets(
    slide,
    x: float,
    y: float,
    w: float,
    h: float,
    items: list[str],
    palette: dict[str, Any],
    *,
    dark: bool = False,
    size: float = 17,
    space_after: float = 8,
    line_spacing: float = 1.15,
):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = box.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.margin_left = frame.margin_right = Inches(0.02)
    frame.margin_top = frame.margin_bottom = Inches(0.02)
    color = palette["white"] if dark else palette["ink"]
    for index, item in enumerate(items or ["보고서의 핵심 근거를 확인해 주세요."]):
        paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
        paragraph.text = f"• {item}"
        paragraph.font.name = "맑은 고딕"
        paragraph.font.size = Pt(size)
        paragraph.font.color.rgb = _rgb(color)
        paragraph.space_after = Pt(space_after)
        paragraph.line_spacing = line_spacing
    return box


def _add_header(slide, item: dict[str, Any], palette: dict[str, Any], *, dark: bool = False):
    color = palette["white"] if dark else palette["ink"]
    muted = palette["pale"] if dark else palette["muted"]
    variant = str(palette.get("variant_id") or "editorial_axis")
    if variant == "field_brief":
        _shape(slide, 0.72, 0.43, 0.42, 0.10, palette[item["accent"]])
    elif variant == "evidence_signal":
        _shape(slide, 0.72, 1.66, 11.88, 0.025, palette["line"])
    elif variant == "cobalt_report":
        _shape(slide, 0.45, 0.43, 0.055, 1.18, palette[item["accent"]])
    _text(slide, 0.72, 0.45, 5.7, 0.28, item["eyebrow"] or f"EVALUATION · {item['slide_number']:02d}", size=10.5, color=muted, bold=True)
    _text(slide, 0.72, 0.82, 11.8, 0.78, item["title"], size=35, color=color, bold=True)


def _add_footer(slide, number: int, palette: dict[str, str], *, dark: bool = False):
    color = palette["pale"] if dark else palette["muted"]
    _text(slide, 0.72, 7.10, 4.6, 0.18, "KODAME · ODA 종료평가 브리핑", size=9, color=color)
    _text(slide, 12.05, 7.06, 0.55, 0.22, f"{number:02d}", size=10, color=color, bold=True, align=PP_ALIGN.RIGHT)


def _add_notes(slide, item: dict[str, Any], project: dict[str, Any]):
    sections = ", ".join(item.get("source_sections") or []) or "해당 슬라이드에 대응하는 평가보고서 섹션"
    documents = ", ".join(item.get("source_documents") or []) or "연결된 보고서 근거자료"
    try:
        slide.notes_slide.notes_text_frame.text = (
            f"{item.get('speaker_notes') or item.get('body') or item.get('headline') or ''}\n\n"
            "[Sources]\n"
            f"- KODAME 평가보고서: {project.get('title') or '현재 사업'}\n"
            f"- 근거 섹션: {sections}\n"
            f"- 근거 문서: {documents}\n"
            "- 수치와 판단은 저장된 최신 평가결과 및 보고서 본문에서 요약"
        )
    except (AttributeError, ValueError):
        pass


def _render_cover(slide, item, plan, context, p):
    variant = str(p.get("variant_id") or "editorial_axis")
    if variant == "cobalt_report":
        slide.background.fill.solid(); slide.background.fill.fore_color.rgb = _rgb(p["white"])
        _shape(slide, 0, 0, 4.10, 7.5, p["blue"])
        _text(slide, 0.76, 0.68, 2.72, 0.32, item["eyebrow"] or "ODA EVALUATION", size=11, color=p["white"], bold=True)
        _text(slide, 0.76, 2.08, 2.75, 1.05, item["metric_value"] or str(context["overall"].get("score") or "—"), size=46, color=p["white"], bold=True)
        _text(slide, 0.76, 3.16, 2.72, 0.72, item["metric_label"] or "종합 평가결과", size=16, color=p["white"], bold=True)
        _text(slide, 4.72, 1.25, 7.58, 1.28, plan["deck_title"] or "종료평가 결과 브리핑", size=50, color=p["ink"], bold=True)
        _text(slide, 4.74, 2.93, 7.28, 1.40, plan["subtitle"] or context["project"]["title"], size=22, color=p["muted"])
        _text(slide, 4.74, 5.92, 7.2, 0.60, " · ".join(x for x in (context["project"].get("country"), context["project"].get("period"), context["project"].get("budget")) if x), size=14, color=p["muted"])
        _text(slide, 0.76, 7.10, 2.8, 0.18, "KODAME · ODA 평가", size=9, color=p["white"])
        _text(slide, 12.05, 7.06, 0.55, 0.22, "01", size=10, color=p["muted"], bold=True, align=PP_ALIGN.RIGHT)
        return

    slide.background.fill.solid(); slide.background.fill.fore_color.rgb = _rgb(p["navy"])
    _shape(slide, 0, 0, 0.18, 7.5, p[item["accent"]])
    _text(slide, 0.82, 0.62, 4.6, 0.3, item["eyebrow"] or "ODA EVALUATION BRIEFING", size=11, color=p["pale"], bold=True)
    _text(slide, 0.82, 1.45, 9.0, 1.78, plan["deck_title"] or "종료평가 결과 브리핑", size=50, color=p["white"], bold=True)
    _text(slide, 0.84, 3.52, 9.1, 1.20, plan["subtitle"] or context["project"]["title"], size=22, color=p["pale"], bold=False)
    if variant == "evidence_signal":
        _shape(slide, 9.95, 1.38, 2.55, 0.08, p[item["accent"]])
    else:
        _shape(slide, 9.95, 0.58, 2.55, 5.78, p[item["accent"]], radius=variant != "field_brief")
    _text(slide, 10.25, 1.05, 1.95, 0.42, f"{SLIDE_COUNT} SLIDES", size=16, color=p["white"], bold=True)
    _text(slide, 10.25, 2.05, 1.9, 1.35, item["metric_value"] or str(context["overall"].get("score") or "—"), size=46, color=p["white"], bold=True)
    _text(slide, 10.25, 3.43, 1.85, 0.8, item["metric_label"] or "종합 평가결과", size=16, color=p["white"], bold=True)
    _text(slide, 0.84, 6.35, 8.8, 0.5, " · ".join(x for x in (context["project"].get("country"), context["project"].get("period"), context["project"].get("budget")) if x), size=14, color=p["pale"])
    _add_footer(slide, 1, p, dark=True)


def _render_statement(slide, item, p):
    slide.background.fill.solid(); slide.background.fill.fore_color.rgb = _rgb(p["paper"])
    _add_header(slide, item, p)
    _shape(slide, 0.72, 1.86, 7.2, 4.72, p["white"], radius=True, line=p["line"])
    _text(slide, 1.08, 2.22, 6.52, 1.55, item["headline"], size=30, color=p["ink"], bold=True)
    _text(slide, 1.08, 4.08, 6.45, 1.75, item["body"], size=18, color=p["muted"])
    _shape(slide, 8.35, 1.86, 4.25, 2.0, p[item["accent"]], radius=True)
    _text(slide, 8.75, 2.18, 3.45, 0.78, item["metric_value"] or "핵심", size=36, color=p["white"], bold=True)
    _text(slide, 8.75, 3.02, 3.3, 0.5, item["metric_label"], size=16, color=p["white"], bold=True)
    _bullets(slide, 8.48, 4.15, 3.9, 2.15, item["bullets"], p, size=16.5)
    _add_footer(slide, item["slide_number"], p)


def _render_split(slide, item, p):
    slide.background.fill.solid(); slide.background.fill.fore_color.rgb = _rgb(p["white"])
    _add_header(slide, item, p)
    _shape(slide, 0.72, 1.86, 5.83, 4.82, p["pale"], radius=True)
    _shape(slide, 6.78, 1.86, 5.83, 4.82, p["paper"], radius=True, line=p["line"])
    _text(slide, 1.08, 2.16, 5.0, 1.18, _clean_text(item["headline"], 60) or "사업이 해결하려는 문제", size=22, color=p["ink"], bold=True)
    _bullets(slide, 1.08, 3.56, 5.0, 2.64, item["bullets"][:4], p, size=16)
    right_heading = _clean_text(item.get("metric_label"), 42) or "사업 설계의 대응"
    _text(slide, 7.15, 2.16, 4.96, 0.78, right_heading, size=24, color=p[item["accent"]], bold=True)
    _bullets(slide, 7.15, 3.25, 4.98, 2.95, (item["secondary_bullets"] or item["bullets"])[:4], p, size=16)
    _add_footer(slide, item["slide_number"], p)


def _render_snapshot(slide, item, p):
    slide.background.fill.solid(); slide.background.fill.fore_color.rgb = _rgb(p["paper"])
    _add_header(slide, item, p)
    _text(slide, 0.82, 1.92, 4.28, 1.55, item["headline"], size=27, color=p["ink"], bold=True)
    _text(slide, 0.84, 3.78, 4.2, 1.72, item["body"], size=17, color=p["muted"])
    metrics = item["metrics"][:4] or [{
        "value": item["metric_value"] or "확인",
        "label": item["metric_label"] or "핵심 사업정보",
        "context": "저장된 보고서와 사업자료 기준",
    }]
    for index, metric in enumerate(metrics):
        column, row = index % 2, index // 2
        x = 5.55 + column * 3.53
        y = 1.94 + row * 2.23
        if str(p.get("variant_id") or "") == "field_brief":
            _shape(slide, x, y + 1.96, 3.06, 0.035, p[item["accent"]])
        else:
            _shape(slide, x, y, 3.24, 2.04, p["white"], radius=True, line=p["line"])
        _text(slide, x + 0.28, y + 0.26, 2.67, 0.56, metric["value"], size=25, color=p[item["accent"]], bold=True)
        _text(slide, x + 0.28, y + 0.88, 2.67, 0.42, metric["label"], size=16, color=p["ink"], bold=True)
        _text(slide, x + 0.28, y + 1.30, 2.67, 0.58, metric["context"], size=16, color=p["muted"])
    _add_footer(slide, item["slide_number"], p)


def _render_pdm_flow(slide, item, p):
    slide.background.fill.solid(); slide.background.fill.fore_color.rgb = _rgb(p["paper"])
    _add_header(slide, item, p)
    steps = _ordered_pdm_steps(item)
    colors = [p["blue"], p["teal"], p["orange"], p["navy"]]
    for i, step in enumerate(steps):
        x = 0.72 + i * 3.08
        _shape(slide, x, 2.05, 2.66, 3.52, colors[i], radius=True)
        _text(slide, x + 0.28, 2.37, 2.05, 0.4, f"0{i + 1}", size=16, color=p["white"], bold=True)
        _text(slide, x + 0.28, 3.0, 2.08, 1.58, step, size=21, color=p["white"], bold=True, valign=MSO_ANCHOR.MIDDLE)
        if i < 3:
            _text(slide, x + 2.68, 3.38, 0.38, 0.38, "→", size=23, color=p["muted"], bold=True, align=PP_ALIGN.CENTER)
    _text(slide, 0.92, 5.93, 11.75, 0.66, item["body"], size=16.5, color=p["muted"], align=PP_ALIGN.CENTER)
    _add_footer(slide, item["slide_number"], p)


def _render_timeline(slide, item, p):
    slide.background.fill.solid(); slide.background.fill.fore_color.rgb = _rgb(p["white"])
    _add_header(slide, item, p)
    steps = (item["steps"] or item["bullets"] or ["목적과 범위", "문헌검토", "교차검증", "종합판단"])[:5]
    _shape(slide, 2.05, 2.1, 0.08, 4.18, p["line"])
    for i, step in enumerate(steps):
        y = 2.05 + i * (3.82 / max(1, len(steps) - 1))
        _shape(slide, 1.75, y, 0.68, 0.68, p[item["accent"]], radius=True)
        _text(slide, 1.75, y + 0.10, 0.68, 0.34, str(i + 1), size=16, color=p["white"], bold=True, align=PP_ALIGN.CENTER)
        _text(slide, 2.78, y - 0.02, 5.7, 0.72, step, size=19, color=p["ink"], bold=True, valign=MSO_ANCHOR.MIDDLE)
    _shape(slide, 8.65, 2.05, 3.85, 4.25, p["pale"], radius=True)
    _text(slide, 9.02, 2.34, 3.08, 1.30, _clean_text(item["headline"], 72), size=19, color=p["ink"], bold=True)
    _text(slide, 9.02, 3.84, 3.05, 2.05, _clean_text(item["body"], 210), size=16, color=p["muted"])
    _add_footer(slide, item["slide_number"], p)


def _render_metric_grid(slide, item, p):
    slide.background.fill.solid(); slide.background.fill.fore_color.rgb = _rgb(p["white"])
    _add_header(slide, item, p)
    metrics = item["metrics"][:4] or [{
        "value": item["metric_value"] or "—",
        "label": item["metric_label"] or "핵심 성과",
        "context": "보고서 근거 기준",
    }]
    width = 11.74 / len(metrics)
    for index, metric in enumerate(metrics):
        x = 0.78 + index * width
        _text(slide, x + 0.06, 1.98, width - 0.24, 0.72, metric["value"], size=30, color=p[item["accent"]], bold=True)
        _text(slide, x + 0.06, 2.82, width - 0.24, 0.54, metric["label"], size=17, color=p["ink"], bold=True)
        _text(slide, x + 0.06, 3.44, width - 0.24, 0.68, metric["context"], size=16, color=p["muted"])
        if index < len(metrics) - 1:
            _shape(slide, x + width - 0.08, 2.02, 0.025, 2.08, p["line"])
    _shape(slide, 0.82, 4.52, 11.69, 0.045, p[item["accent"]])
    _bullets(slide, 0.90, 4.92, 7.40, 1.46, item["bullets"], p, size=16)
    _text(slide, 8.72, 4.90, 3.52, 1.24, item["headline"], size=20, color=p["ink"], bold=True, valign=MSO_ANCHOR.MIDDLE)
    _add_footer(slide, item["slide_number"], p)


def _render_score_bars(slide, item, context, p):
    slide.background.fill.solid(); slide.background.fill.fore_color.rgb = _rgb(p["paper"])
    _add_header(slide, item, p)
    criteria = context.get("criteria") or []
    if not criteria:
        criteria = [{"name": label, "currentScore4": 0} for label in ("적절성", "일관성", "효과성", "효율성", "지속가능성")]
    for i, criterion in enumerate(criteria[:5]):
        y = 2.05 + i * 0.83
        score = max(0.0, min(4.0, float(criterion.get("currentScore4") or 0)))
        _text(slide, 0.85, y, 1.45, 0.42, str(criterion.get("name") or "기준"), size=16, color=p["ink"], bold=True)
        _shape(slide, 2.34, y + 0.06, 5.15, 0.25, p["line"], radius=True)
        if criterion.get('currentScore4') is not None:
            _shape(slide, 2.34, y + 0.06, max(0.08, 5.15 * score / 4), 0.25, p[item["accent"]], radius=True)
        _text(slide, 7.68, y - 0.05, 0.75, 0.42, f"{score:.1f}" if criterion.get('currentScore4') is not None else '보류', size=17, color=p[item["accent"]], bold=True, align=PP_ALIGN.RIGHT)
    _shape(slide, 8.72, 1.94, 3.88, 4.65, p["white"], radius=True, line=p["line"])
    _text(slide, 9.08, 2.26, 3.15, 1.42, _clean_text(item["headline"], 78), size=18.5, color=p["ink"], bold=True)
    _bullets(
        slide, 9.06, 3.70, 3.15, 2.56, item["bullets"][:3], p,
        size=16, space_after=4, line_spacing=1.05,
    )
    _text(slide, 0.88, 6.40, 7.42, 0.35, "4점 척도 · 최신 저장 평가결과 기준", size=10.5, color=p["muted"])
    _add_footer(slide, item["slide_number"], p)


def _render_comparison(slide, item, p):
    slide.background.fill.solid(); slide.background.fill.fore_color.rgb = _rgb(p["white"])
    _add_header(slide, item, p)
    _text(slide, 0.82, 1.86, 5.65, 0.55, "확인된 강점", size=24, color=p["teal"], bold=True)
    _shape(slide, 0.82, 2.55, 5.55, 3.82, p["pale"], radius=True)
    _bullets(slide, 1.16, 2.91, 4.85, 2.95, item["bullets"], p, size=17)
    _text(slide, 6.83, 1.86, 5.55, 0.55, "보완이 필요한 지점", size=24, color=p["orange"], bold=True)
    _shape(slide, 6.83, 2.55, 5.55, 3.82, p["sand"], radius=True)
    _bullets(slide, 7.17, 2.91, 4.85, 2.95, item["secondary_bullets"] or item["bullets"], p, size=17)
    _add_footer(slide, item["slide_number"], p)


def _render_evidence(slide, item, p):
    slide.background.fill.solid(); slide.background.fill.fore_color.rgb = _rgb(p["navy"])
    _add_header(slide, item, p, dark=True)
    _text(slide, 0.94, 1.95, 10.85, 1.72, f"“{item['quote'] or item['headline']}”", size=29, color=p["white"], bold=True)
    _shape(slide, 0.94, 4.15, 11.55, 0.06, p[item["accent"]])
    _bullets(slide, 0.98, 4.65, 7.4, 1.65, item["bullets"], p, dark=True, size=17)
    _shape(slide, 9.02, 4.48, 3.45, 1.68, p[item["accent"]], radius=True)
    _text(slide, 9.35, 4.75, 2.75, 0.52, item["metric_value"] or "교훈", size=26, color=p["white"], bold=True)
    _text(slide, 9.35, 5.36, 2.72, 0.48, item["metric_label"], size=16, color=p["white"], bold=True)
    _add_footer(slide, item["slide_number"], p, dark=True)


def _render_roadmap(slide, item, p):
    slide.background.fill.solid(); slide.background.fill.fore_color.rgb = _rgb(p["paper"])
    _add_header(slide, item, p)
    steps = (item["steps"] or item["bullets"] or ["즉시 보완", "운영 내재화", "성과 확산"])[:4]
    for i, step in enumerate(steps):
        column, row = i % 2, i // 2
        x = 0.78 + column * 6.02
        y = 1.90 + row * 2.10
        title, details = _roadmap_step_parts(step)
        _shape(slide, x, y, 5.74, 1.82, p["white"], radius=True, line=p["line"])
        _text(slide, x + 0.25, y + 0.22, 0.54, 0.32, f"P{i + 1}", size=16, color=p[item["accent"]], bold=True)
        _text(slide, x + 0.96, y + 0.18, 4.44, 0.54, title, size=20, color=p["ink"], bold=True)
        _text(slide, x + 0.96, y + 0.83, 4.44, 0.70, "  ·  ".join(details), size=16, color=p["muted"])
    _text(slide, 0.82, 6.20, 11.6, 0.48, _clean_text(item["body"], 150), size=16, color=p["muted"], align=PP_ALIGN.CENTER)
    _add_footer(slide, item["slide_number"], p)


def _render_closing(slide, item, p):
    slide.background.fill.solid(); slide.background.fill.fore_color.rgb = _rgb(p["white"])
    _shape(slide, 0, 0, 4.15, 7.5, p[item["accent"]])
    _text(slide, 0.72, 0.68, 2.9, 0.3, item["eyebrow"] or "DECISION", size=11, color=p["white"], bold=True)
    _text(slide, 0.72, 1.64, 2.86, 2.34, item["metric_value"] or "NEXT", size=43, color=p["white"], bold=True, valign=MSO_ANCHOR.MIDDLE)
    _text(slide, 4.75, 0.88, 7.78, 1.55, item["headline"] or item["title"], size=35, color=p["ink"], bold=True)
    _text(slide, 4.78, 2.72, 7.15, 1.0, item["body"], size=18, color=p["muted"])
    _bullets(slide, 4.78, 4.05, 7.2, 2.20, item["bullets"], p, size=18, space_after=6, line_spacing=1.08)
    _text(slide, 0.72, 7.10, 3.05, 0.18, "KODAME · ODA 종료평가 브리핑", size=9, color=p["white"])
    _text(slide, 12.05, 7.06, 0.55, 0.22, f"{item['slide_number']:02d}", size=10, color=p["muted"], bold=True, align=PP_ALIGN.RIGHT)


def build_presentation_bytes(plan: dict[str, Any], context: dict[str, Any]) -> bytes:
    presentation = Presentation()
    presentation.slide_width = Inches(13.333333)
    presentation.slide_height = Inches(7.5)
    blank = presentation.slide_layouts[6]
    palette: dict[str, Any] = dict(PALETTES[plan["design"]["palette"]])
    palette["variant_id"] = plan["design"].get("variant_id") or "editorial_axis"
    renderers = {
        "statement": lambda s, i: _render_statement(s, i, palette),
        "split": lambda s, i: _render_split(s, i, palette),
        "snapshot": lambda s, i: _render_snapshot(s, i, palette),
        "pdm_flow": lambda s, i: _render_pdm_flow(s, i, palette),
        "timeline": lambda s, i: _render_timeline(s, i, palette),
        "metric_grid": lambda s, i: _render_metric_grid(s, i, palette),
        "score_bars": lambda s, i: _render_score_bars(s, i, context, palette),
        "comparison": lambda s, i: _render_comparison(s, i, palette),
        "evidence": lambda s, i: _render_evidence(s, i, palette),
        "roadmap": lambda s, i: _render_roadmap(s, i, palette),
        "closing": lambda s, i: _render_closing(s, i, palette),
    }
    for item in plan["slides"]:
        slide = presentation.slides.add_slide(blank)
        if item["layout"] == "cover":
            _render_cover(slide, item, plan, context, palette)
        else:
            renderers.get(item["layout"], renderers[LAYOUT_SEQUENCE[item["slide_number"] - 1]])(slide, item)
        _add_notes(slide, item, context["project"])
    output = BytesIO()
    presentation.save(output)
    return output.getvalue()


def _update(export_id: uuid.UUID, progress: int, stage: str, message: str) -> None:
    with connection() as conn, conn.transaction():
        conn.execute(
            """UPDATE presentation_exports
                  SET status='running',progress=%s,stage=%s,message=%s,
                      started_at=COALESCE(started_at,now()),updated_at=now()
                WHERE id=%s""",
            (progress, stage, message, export_id),
        )


def run_presentation_export(export_id: uuid.UUID, project_id: uuid.UUID | None = None) -> None:
    with tenant_context(project_id, system=project_id is None):
        try:
            _update(export_id, 8, "collecting", "저장된 27개 보고서 섹션과 최신 평가결과를 읽는 중")
            context, source, sensitive_flags = _collect_presentation_source()
            qa_feedback: list[str] = []
            attempt_history: list[dict[str, Any]] = []
            data: bytes | None = None
            plan: dict[str, Any] | None = None
            validation: dict[str, Any] | None = None
            quality_report: dict[str, Any] | None = None
            max_attempts = min(
                PRESENTATION_MAX_GENERATION_ATTEMPTS,
                max(1, int(build_reference_context(source, str(export_id))["max_generation_attempts"])),
            )
            for attempt in range(1, max_attempts + 1):
                reference_context = build_reference_context(source, f"{export_id}:{attempt}")
                reference_context["quality_target"] = max(
                    float(reference_context.get("quality_target") or 90),
                    PRESENTATION_MIN_QUALITY_SCORE,
                )
                _update(
                    export_id,
                    18 + attempt * 8,
                    "planning",
                    f"Claude 고추론으로 참조 양식 기반 발표 설계를 구성하는 중 ({attempt}/{max_attempts})",
                )
                try:
                    raw_plan = request_presentation_plan(source, reference_context, qa_feedback)
                    candidate_plan = normalize_presentation_plan(
                        raw_plan, context["project"], source, reference_context
                    )
                    _update(
                        export_id,
                        48 + attempt * 6,
                        "rendering",
                        f"통제된 디자인 변형으로 편집 가능한 PPTX를 조판하는 중 ({attempt}/{max_attempts})",
                    )
                    candidate_data = build_presentation_bytes(candidate_plan, context)
                    structural = validate_presentation_bytes(candidate_data, candidate_plan)
                    _update(
                        export_id,
                        68 + attempt * 5,
                        "visual_qa",
                        f"전체 슬라이드를 실제 렌더링하고 100점 기준으로 평가하는 중 ({attempt}/{max_attempts})",
                    )
                    rendered = render_and_validate_presentation(candidate_data, candidate_plan)
                    candidate_quality = score_presentation_quality(
                        candidate_data,
                        candidate_plan,
                        structural,
                        rendered,
                        reference_context,
                    )
                    attempt_history.append({
                        "attempt": attempt,
                        "variant_id": candidate_plan["design"]["variant_id"],
                        "variant_seed": candidate_plan["design"]["variant_seed"],
                        "quality_score": candidate_quality["quality_score"],
                        "issues": candidate_quality["issues"],
                    })
                    if candidate_quality["quality_score"] >= reference_context["quality_target"] and candidate_quality["passed"]:
                        data = candidate_data
                        plan = candidate_plan
                        validation = {**structural, **rendered}
                        quality_report = candidate_quality
                        break
                    qa_feedback = candidate_quality["issues"] or [
                        f"품질점수 {candidate_quality['quality_score']}점으로 목표 {reference_context['quality_target']}점 미달"
                    ]
                except Exception as attempt_error:
                    qa_feedback = [f"이전 생성본 검증 실패: {str(attempt_error)[:700]}"]
                    attempt_history.append({
                        "attempt": attempt,
                        "quality_score": 0,
                        "issues": qa_feedback,
                    })
            if data is None or plan is None or validation is None or quality_report is None:
                raise RuntimeError(
                    f"발표자료가 {max_attempts}회 내 품질 {PRESENTATION_MIN_QUALITY_SCORE:.0f}점 기준을 통과하지 못했습니다: "
                    f"{qa_feedback}"
                )
            validation["quality"] = quality_report
            validation["quality_attempts"] = attempt_history
            validation["design"] = plan["design"]
            validation["local_sensitive_flags"] = sensitive_flags
            validation["source_section_coverage"] = len({
                section for slide in plan["slides"] for section in slide.get("source_sections", [])
            })
            validation["source_document_count"] = len({
                document for slide in plan["slides"] for document in slide.get("source_documents", [])
            })
            validation["model"] = current_llm_model()
            validation["reasoning_effort"] = "high"
            _update(
                export_id, 94, "saving",
                f"품질 {quality_report['quality_score']:.0f}점으로 통과한 {SLIDE_COUNT}장 발표자료를 저장하는 중",
            )
            PRESENTATION_DIR.mkdir(parents=True, exist_ok=True)
            project_name = re.sub(r"[^0-9A-Za-z가-힣._-]+", "_", context["project"]["title"]).strip("._")[:70] or "ODA_사업"
            file_name = f"{project_name}_평가결과_고품질_발표자료_{SLIDE_COUNT}장.pptx"
            output_path = PRESENTATION_DIR / f"{export_id}.pptx"
            output_path.write_bytes(data)
            with connection() as conn, conn.transaction():
                conn.execute(
                    """UPDATE presentation_exports
                          SET status='completed',progress=100,stage='completed',
                              message=%s,output_path=%s,file_name=%s,
                              validation=%s,error_message=NULL,completed_at=now(),updated_at=now()
                        WHERE id=%s""",
                    (
                        f"Claude 고추론·참조 양식·렌더 QA {quality_report['quality_score']:.0f}점 발표자료 생성 완료",
                        str(output_path), file_name, Jsonb(validation), export_id,
                    ),
                )
        except Exception as exc:
            with connection() as conn, conn.transaction():
                conn.execute(
                    """UPDATE presentation_exports
                          SET status='failed',stage='failed',message='발표자료 생성 실패',
                              error_message=%s,completed_at=now(),updated_at=now()
                        WHERE id=%s""",
                    (str(exc)[:2000], export_id),
                )
