from __future__ import annotations

import math
import re
import subprocess
import tempfile
from io import BytesIO
from pathlib import Path
from typing import Any

from pptx import Presentation
from pypdf import PdfReader

from .presentation_prompt import SLIDE_COUNT


EMU_PER_INCH = 914400
SLIDE_WIDTH_IN = 13.333333
SLIDE_HEIGHT_IN = 7.5


def _inches(value: int) -> float:
    return float(value) / EMU_PER_INCH


def _shape_font_sizes(shape) -> list[float]:
    sizes: list[float] = []
    if not getattr(shape, "has_text_frame", False):
        return sizes
    for paragraph in shape.text_frame.paragraphs:
        if paragraph.font.size:
            sizes.append(float(paragraph.font.size.pt))
        for run in paragraph.runs:
            if run.font.size:
                sizes.append(float(run.font.size.pt))
    return sizes


def _estimated_lines(text: str, width_in: float, font_size: float) -> int:
    if not text.strip():
        return 0
    # A mixed Korean/Latin presentation averages a little over half an em per
    # character after actual wrapping.  Layout-specific renderers carry the
    # stricter content budgets; this generic check catches gross overflow
    # without rejecting valid display typography.
    chars_per_line = max(5, int(width_in * 72 / max(8.0, font_size * 0.58)))
    return sum(max(1, math.ceil(len(line.strip()) / chars_per_line)) for line in text.splitlines() or [text])


def _boxes_overlap(left: Any, right: Any, tolerance: float = 0.04) -> bool:
    lx, ly, lw, lh = (_inches(left.left), _inches(left.top), _inches(left.width), _inches(left.height))
    rx, ry, rw, rh = (_inches(right.left), _inches(right.top), _inches(right.width), _inches(right.height))
    overlap_w = min(lx + lw, rx + rw) - max(lx, rx)
    overlap_h = min(ly + lh, ry + rh) - max(ly, ry)
    return overlap_w > tolerance and overlap_h > tolerance


def validate_presentation_bytes(data: bytes, plan: dict[str, Any]) -> dict[str, Any]:
    presentation = Presentation(BytesIO(data))
    if len(presentation.slides) != SLIDE_COUNT:
        raise RuntimeError(f"PPTX 슬라이드 수가 {SLIDE_COUNT}장이 아닙니다.")

    empty_slides: list[int] = []
    out_of_bounds: list[str] = []
    possible_overflow: list[str] = []
    text_overlaps: list[str] = []
    undersized_body: list[str] = []
    slide_texts: list[str] = []

    for slide_index, slide in enumerate(presentation.slides, 1):
        text_shapes = []
        visible_parts = []
        for shape_index, shape in enumerate(slide.shapes, 1):
            x, y = _inches(shape.left), _inches(shape.top)
            width, height = _inches(shape.width), _inches(shape.height)
            if x < -0.01 or y < -0.01 or x + width > SLIDE_WIDTH_IN + 0.01 or y + height > SLIDE_HEIGHT_IN + 0.01:
                out_of_bounds.append(f"{slide_index}:{shape_index}")
            text = str(getattr(shape, "text", "") or "").strip()
            if not text:
                continue
            visible_parts.append(text)
            text_shapes.append((shape_index, shape))
            sizes = _shape_font_sizes(shape)
            font_size = min(sizes) if sizes else 18.0
            is_metadata = y < 0.78 or (y >= 6.18 and height <= 0.72)
            if not is_metadata and font_size < 15.9:
                undersized_body.append(f"{slide_index}:{shape_index}:{font_size:.1f}pt")
            capacity = max(1, int(height * 72 / max(10.0, font_size * 1.18)))
            paragraph_spacing = 0.0
            for paragraph in getattr(shape.text_frame, "paragraphs", []):
                if paragraph.space_before:
                    paragraph_spacing += float(paragraph.space_before.pt)
                if paragraph.space_after:
                    paragraph_spacing += float(paragraph.space_after.pt)
            spacing_lines = math.ceil(paragraph_spacing / max(10.0, font_size * 1.18))
            if _estimated_lines(text, width, font_size) + spacing_lines > capacity + 1:
                possible_overflow.append(f"{slide_index}:{shape_index}")
        visible = " ".join(visible_parts).strip()
        slide_texts.append(visible)
        if len(visible) < 20:
            empty_slides.append(slide_index)
        for left_index in range(len(text_shapes)):
            left_no, left_shape = text_shapes[left_index]
            for right_no, right_shape in text_shapes[left_index + 1:]:
                if _boxes_overlap(left_shape, right_shape):
                    text_overlaps.append(f"{slide_index}:{left_no}-{right_no}")

    missing_titles = [
        item["slide_number"]
        for item, slide_text in zip(plan["slides"], slide_texts)
        if item.get("layout") not in {"cover", "closing"}
        and str(item.get("title") or "").strip()
        and str(item["title"]).strip() not in slide_text
    ]
    errors = {
        "empty_slides": empty_slides,
        "out_of_bounds": out_of_bounds,
        "possible_text_overflow": possible_overflow,
        "text_box_overlaps": text_overlaps,
        "undersized_body_text": undersized_body,
        "missing_titles": missing_titles,
    }
    failed = {key: value for key, value in errors.items() if value}
    if failed:
        raise RuntimeError(f"PPTX 레이아웃 안전성 검증 실패: {failed}")
    return {
        "slide_count": len(presentation.slides),
        "file_size": len(data),
        "layouts": [item["layout"] for item in plan["slides"]],
        "unique_layouts": len({item["layout"] for item in plan["slides"]}),
        **errors,
    }


def missing_rendered_content(required: list[str], page_readings: list[str]) -> list[str]:
    """Accept a complete phrase in either reading order of the same PDF page.

    Poppler can interleave adjacent table cells when a cell wraps. pypdf's
    content-stream order preserves those cells. Do not join readings or match
    individual words: that could conceal a genuinely missing sentence.
    """
    readings = [re.sub(r"\s+", "", text) for text in page_readings]
    return [text for text in required
            if not any(re.sub(r"\s+", "", text) in visible for visible in readings)]


def render_and_validate_presentation(data: bytes, plan: dict[str, Any]) -> dict[str, Any]:
    """Render with LibreOffice/Poppler and fail before delivery on broken output."""
    expected_count = plan.get("slide_count", SLIDE_COUNT)
    with tempfile.TemporaryDirectory(prefix="kodame_presentation_qa_") as directory:
        root = Path(directory)
        pptx_path = root / "presentation.pptx"
        pptx_path.write_bytes(data)
        converted = subprocess.run(
            ["soffice", "--headless", "--convert-to", "pdf", "--outdir", str(root), str(pptx_path)],
            capture_output=True,
            text=True,
            timeout=180,
        )
        pdf_path = root / "presentation.pdf"
        if converted.returncode != 0 or not pdf_path.is_file():
            raise RuntimeError(f"PPTX 렌더링 검증 실패: {(converted.stderr or converted.stdout)[-500:]}")
        rendered = subprocess.run(
            ["pdftoppm", "-png", "-r", "110", str(pdf_path), str(root / "slide")],
            capture_output=True,
            text=True,
            timeout=180,
        )
        if rendered.returncode != 0:
            raise RuntimeError(f"PPTX 슬라이드 이미지 검증 실패: {rendered.stderr[-500:]}")
        images = sorted(root.glob("slide-*.png"))
        if len(images) != expected_count or any(path.stat().st_size < 5000 for path in images):
            raise RuntimeError(f"PPTX 렌더링 페이지 완전성 실패: {len(images)}/{expected_count}")
        extracted = subprocess.run(
            ["pdftotext", str(pdf_path), "-"],
            capture_output=True,
            text=True,
            timeout=60,
        )
        if extracted.returncode != 0:
            raise RuntimeError(f"PPTX 렌더링 텍스트 추출 실패: {extracted.stderr[-500:]}")
        pages = (extracted.stdout or "").split("\f")
        stream_pages = [page.extract_text() or "" for page in PdfReader(pdf_path).pages]
        readings = [[pages[i] if i < len(pages) else "", stream_pages[i]]
                    for i in range(len(stream_pages))]
        if len(readings) != expected_count:
            raise RuntimeError(f"PPTX 텍스트 페이지 완전성 실패: {len(readings)}/{expected_count}")
        missing_rendered_titles = [
            item["slide_number"]
            for index, item in enumerate(plan["slides"])
            if item.get("layout") not in {"cover", "closing"}
            and missing_rendered_content([str(item.get("title") or "")], readings[index])
        ]
        if missing_rendered_titles:
            raise RuntimeError(f"PPTX 렌더링 텍스트 검증 실패: 제목 누락 {missing_rendered_titles}")
        if plan.get("slide_count") in (15, 30):
            for index, item in enumerate(plan["slides"]):
                required = [cell for row in item.get("rows", []) for cell in row]
                required += [b.get("text", "") for b in item.get("blocks", [])]
                # Cover/TOC content is supplied by the profile, not model blocks.
                if item.get("layout") in ("cover", "toc"):
                    continue
                missing = [s[:50] for s in missing_rendered_content(required, readings[index])]
                if missing:
                    raise RuntimeError(f"{index+1}페이지 실제 렌더링에서 본문 누락: {missing[:3]}")
        return {
            "render_engine": "LibreOffice Impress + Poppler",
            "text_readers": ["Poppler", "pypdf content-stream order"],
            "rendered_slide_count": len(images),
            "rendered_titles_verified": expected_count,
            "smallest_rendered_png_bytes": min(path.stat().st_size for path in images),
        }


def score_presentation_quality(
    data: bytes,
    plan: dict[str, Any],
    structural: dict[str, Any],
    rendered: dict[str, Any],
    reference_context: dict[str, Any],
) -> dict[str, Any]:
    """Score the deliverable against an explicit 100-point editorial rubric."""
    presentation = Presentation(BytesIO(data))
    issues: list[str] = []
    roles = [str(item.get("role") or "") for item in plan.get("slides") or []]
    expected_roles = [
        "cover", "executive_summary", "context", "project_profile", "pdm", "methodology",
        "achievement", "assessment", "crosscutting", "factors_lessons", "recommendations", "closing",
    ]

    narrative = 20.0
    if roles != expected_roles:
        narrative -= 8
        issues.append("12장 서사 역할 순서가 표준 흐름과 다름")
    if any(not str(item.get("title") or "").strip() for item in plan["slides"]):
        narrative -= 4
        issues.append("판단형 제목이 비어 있는 슬라이드가 있음")
    if any(len(str(item.get("title") or "")) > 32 for item in plan["slides"][1:]):
        narrative -= 3
        issues.append("한 줄 제목 예산을 초과한 슬라이드가 있음")
    if any(not (item.get("headline") or item.get("body") or item.get("bullets")) for item in plan["slides"][1:]):
        narrative -= 5
        issues.append("핵심 주장이나 근거가 없는 슬라이드가 있음")

    reference_alignment = 15.0
    design = plan.get("design") or {}
    selected_variant = reference_context.get("selected_variant") or {}
    if not design.get("reference_profile_version"):
        reference_alignment -= 5
        issues.append("참조 프로필 버전이 기록되지 않음")
    if design.get("variant_id") != selected_variant.get("id"):
        reference_alignment -= 4
        issues.append("선택된 디자인 변형이 렌더 계획과 일치하지 않음")
    if design.get("palette") != selected_variant.get("palette"):
        reference_alignment -= 3
        issues.append("참조 변형의 팔레트가 반영되지 않음")
    if structural.get("unique_layouts", 0) < 9:
        reference_alignment -= 3
        issues.append("장표 실루엣 다양성이 9종 미만임")

    hierarchy = 20.0
    font_sizes = [
        size
        for slide in presentation.slides
        for shape in slide.shapes
        for size in _shape_font_sizes(shape)
    ]
    if structural.get("undersized_body_text"):
        hierarchy -= 8
        issues.append("16pt 미만 본문이 있음")
    if min(font_sizes or [0]) <= 0:
        hierarchy -= 4
        issues.append("글꼴 크기를 확인할 수 없는 텍스트가 있음")
    if any(len(item.get("bullets") or []) > 4 for item in plan["slides"]):
        hierarchy -= 4
        issues.append("한 장의 핵심 목록이 4개를 초과함")
    if len(str(plan.get("deck_title") or "")) > 22:
        hierarchy -= 4
        issues.append("표지 제목이 의도한 한 줄 예산을 초과함")

    layout_integrity = 25.0
    for key in (
        "empty_slides", "out_of_bounds", "possible_text_overflow", "text_box_overlaps",
        "undersized_body_text", "missing_titles",
    ):
        if structural.get(key):
            layout_integrity -= 8
            issues.append(f"구조 검증 문제: {key}")
    if rendered.get("rendered_slide_count") != SLIDE_COUNT:
        layout_integrity -= 12
        issues.append("실제 렌더링 장수가 12장이 아님")
    if rendered.get("rendered_titles_verified") != SLIDE_COUNT:
        layout_integrity -= 8
        issues.append("실제 렌더링에서 제목이 누락됨")
    layout_integrity = max(0.0, layout_integrity)

    traceability = 10.0
    for slide_number, (slide, item) in enumerate(zip(presentation.slides, plan["slides"]), 1):
        notes = str(slide.notes_slide.notes_text_frame.text or "")
        if "[Sources]" not in notes:
            traceability -= 2
            issues.append(f"{slide_number}장 [Sources] 발표자 노트 누락")
        if slide_number > 1 and not item.get("source_sections"):
            traceability -= 1
            issues.append(f"{slide_number}장 보고서 근거 섹션 누락")
    traceability = max(0.0, traceability)

    variety = 10.0
    layouts = [str(item.get("layout") or "") for item in plan["slides"]]
    if any(left == right for left, right in zip(layouts, layouts[1:])):
        variety -= 5
        issues.append("인접 슬라이드의 실루엣이 반복됨")
    accents = {str(item.get("accent") or "") for item in plan["slides"]}
    if len(accents) < 2:
        variety -= 3
        issues.append("강조색 운용이 단조로움")
    if not design.get("variant_seed"):
        variety -= 2
        issues.append("재현 가능한 디자인 시드가 기록되지 않음")

    breakdown = {
        "narrative_and_decision_flow": round(max(0.0, narrative), 1),
        "reference_format_alignment": round(max(0.0, reference_alignment), 1),
        "visual_hierarchy_and_density": round(max(0.0, hierarchy), 1),
        "layout_and_render_integrity": round(layout_integrity, 1),
        "source_traceability": round(traceability, 1),
        "controlled_design_variety": round(max(0.0, variety), 1),
    }
    score = round(sum(breakdown.values()), 1)
    return {
        "quality_score": score,
        "quality_target": float(reference_context.get("quality_target") or 90),
        "passed": score >= float(reference_context.get("quality_target") or 90) and not issues,
        "breakdown": breakdown,
        "issues": list(dict.fromkeys(issues)),
    }
