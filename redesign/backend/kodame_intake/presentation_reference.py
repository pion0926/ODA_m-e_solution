from __future__ import annotations

import hashlib
import json
import glob
from collections import Counter
from pathlib import Path
from typing import Any

from pptx import Presentation

from .settings import PRESENTATION_REFERENCE_PROFILE_PATH


DEFAULT_REFERENCE_FOLDER_URL = (
    "https://drive.google.com/drive/folders/"
    "13UG78hLEcJMfWNQo0D5Nc5vMZbtvZahr?usp=drive_link"
)


def load_presentation_reference_profile(path: Path | None = None) -> dict[str, Any]:
    """Load the editable, format-only presentation reference contract."""
    profile_path = path or PRESENTATION_REFERENCE_PROFILE_PATH
    try:
        profile = json.loads(profile_path.read_text(encoding="utf-8"))
    except (OSError, json.JSONDecodeError) as exc:
        raise RuntimeError(f"발표자료 참조 프로필을 읽을 수 없습니다: {profile_path}") from exc
    if not isinstance(profile, dict):
        raise RuntimeError("발표자료 참조 프로필은 JSON 객체여야 합니다.")
    folder = profile.get("reference_folder")
    if not isinstance(folder, dict) or not str(folder.get("url") or "").startswith("https://drive.google.com/"):
        raise RuntimeError("발표자료 참조 프로필의 Google Drive 폴더 URL이 유효하지 않습니다.")
    variants = profile.get("variant_families")
    if not isinstance(variants, list) or len(variants) < 3:
        raise RuntimeError("발표자료 참조 프로필에는 최소 3개 디자인 변형이 필요합니다.")
    return profile


def _seed_material(source: dict[str, Any], generation_key: str) -> str:
    project = (source.get("summary") or {}).get("project") or {}
    return "|".join(
        [
            str(project.get("title") or ""),
            str(project.get("country") or ""),
            str(generation_key or "default"),
        ]
    )


def select_reference_variant(
    profile: dict[str, Any], source: dict[str, Any], generation_key: str = ""
) -> dict[str, Any]:
    """Select a reproducible variant; new export ids yield controlled diversity."""
    variants = profile["variant_families"]
    digest = hashlib.sha256(_seed_material(source, generation_key).encode("utf-8")).digest()
    selected = dict(variants[int.from_bytes(digest[:4], "big") % len(variants)])
    selected["seed"] = digest.hex()[:12]
    return selected


def _font_names(shape: Any) -> list[str]:
    if not getattr(shape, "has_text_frame", False):
        return []
    names: list[str] = []
    for paragraph in shape.text_frame.paragraphs:
        if paragraph.font.name:
            names.append(str(paragraph.font.name))
        for run in paragraph.runs:
            if run.font.name:
                names.append(str(run.font.name))
    return names


def analyze_local_reference_decks(profile: dict[str, Any]) -> list[dict[str, Any]]:
    """Extract format metrics only; no sample text is returned or sent to the model."""
    paths = [str(item.get("path") or "") for item in profile.get("local_reference_decks") or []]
    for pattern in profile.get("local_reference_globs") or []:
        paths.extend(glob.glob(str(pattern), recursive=True))
    metrics: list[dict[str, Any]] = []
    for value in dict.fromkeys(path for path in paths if path):
        path = Path(value)
        if not path.is_file() or path.suffix.lower() != ".pptx":
            continue
        try:
            deck = Presentation(str(path))
        except (OSError, ValueError, KeyError):
            continue
        shape_counts: list[int] = []
        text_box_counts: list[int] = []
        table_counts: list[int] = []
        image_counts: list[int] = []
        fonts: Counter[str] = Counter()
        for slide in deck.slides:
            shapes = list(slide.shapes)
            shape_counts.append(len(shapes))
            text_box_counts.append(sum(1 for shape in shapes if getattr(shape, "has_text_frame", False)))
            table_counts.append(sum(1 for shape in shapes if getattr(shape, "has_table", False)))
            image_counts.append(sum(1 for shape in shapes if getattr(shape, "shape_type", None) == 13))
            fonts.update(name for shape in shapes for name in _font_names(shape))
        slide_count = len(deck.slides)
        metrics.append({
            "file_name": path.name,
            "slide_count": slide_count,
            "slide_ratio": round(float(deck.slide_width) / max(1.0, float(deck.slide_height)), 3),
            "average_shapes_per_slide": round(sum(shape_counts) / max(1, slide_count), 1),
            "average_text_boxes_per_slide": round(sum(text_box_counts) / max(1, slide_count), 1),
            "table_slide_count": sum(1 for value in table_counts if value),
            "image_slide_count": sum(1 for value in image_counts if value),
            "dominant_fonts": [name for name, _count in fonts.most_common(4)],
            "content_excluded": True,
        })
    return metrics


def build_reference_context(
    source: dict[str, Any], generation_key: str = "", *, profile: dict[str, Any] | None = None
) -> dict[str, Any]:
    reference = profile or load_presentation_reference_profile()
    return {
        "profile_version": reference.get("profile_version"),
        "reference_folder": reference.get("reference_folder"),
        "content_firewall": reference.get("content_firewall"),
        "narrative_patterns": reference.get("narrative_patterns") or [],
        "visual_patterns": reference.get("visual_patterns") or [],
        "selected_variant": select_reference_variant(reference, source, generation_key),
        "sample_format_metrics": analyze_local_reference_decks(reference),
        "quality_target": int(reference.get("quality_target") or 90),
        "max_generation_attempts": max(1, min(5, int(reference.get("max_generation_attempts") or 3))),
    }


def prompt_reference_payload(context: dict[str, Any]) -> dict[str, Any]:
    """Return only structure and design cues; never expose reference deck content."""
    return {
        "profile_version": context.get("profile_version"),
        "source_folder_url": (context.get("reference_folder") or {}).get("url")
        or DEFAULT_REFERENCE_FOLDER_URL,
        "content_firewall": context.get("content_firewall"),
        "narrative_patterns": context.get("narrative_patterns"),
        "visual_patterns": context.get("visual_patterns"),
        "selected_variant": context.get("selected_variant"),
        "sample_format_metrics": context.get("sample_format_metrics") or [],
    }
