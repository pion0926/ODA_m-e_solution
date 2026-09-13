from __future__ import annotations

import csv
import io
import re
import struct
import subprocess
import tempfile
import zipfile
import zlib
from contextvars import ContextVar
from pathlib import Path

import olefile
from defusedxml import ElementTree
from docx import Document
from openpyxl import load_workbook
from pypdf import PdfReader
from pptx import Presentation

from .settings import MAX_EXTRACTED_CHARS

MAX_ARCHIVE_ENTRIES = 5000
MAX_ARCHIVE_UNCOMPRESSED = 512 * 1024 * 1024
_full_text = ContextVar("parse_full_text", default=False)

class ParseError(RuntimeError):
    pass

def _trim(text: str) -> str:
    text = text.replace("\x00", "")
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text).strip()
    return text if _full_text.get() else text[:MAX_EXTRACTED_CHARS]

def _safe_zip(path: Path) -> zipfile.ZipFile:
    archive = zipfile.ZipFile(path)
    infos = archive.infolist()
    if len(infos) > MAX_ARCHIVE_ENTRIES or sum(item.file_size for item in infos) > MAX_ARCHIVE_UNCOMPRESSED:
        archive.close()
        raise ParseError("압축 문서의 내부 크기 또는 항목 수가 안전 제한을 초과했습니다.")
    if any(".." in Path(item.filename).parts for item in infos):
        archive.close()
        raise ParseError("압축 문서에 안전하지 않은 경로가 있습니다.")
    return archive

def _pdf(path: Path) -> str:
    reader = PdfReader(str(path), strict=False)
    if reader.is_encrypted:
        try:
            reader.decrypt("")
        except Exception as exc:
            raise ParseError("암호화된 PDF는 처리할 수 없습니다.") from exc
    text = "\n\n".join((page.extract_text() or "") for page in reader.pages)
    if len(text.strip()) >= 50:
        return text
    return _pdf_ocr(path, len(reader.pages) if _full_text.get() else min(len(reader.pages), 100))

def _pdf_ocr(path: Path, page_count: int) -> str:
    if page_count < 1:
        raise ParseError("PDF에 페이지가 없습니다.")
    chunks: list[str] = []
    with tempfile.TemporaryDirectory(prefix="kodame-ocr-") as temp_dir:
        prefix = str(Path(temp_dir) / "page")
        try:
            subprocess.run(
                ["pdftoppm", "-jpeg", "-r", "160", "-f", "1", "-l", str(page_count), str(path), prefix],
                check=True, capture_output=True, timeout=max(180, page_count * 20),
            )
            images = sorted(Path(temp_dir).glob("page-*.jpg"))
            for number, image in enumerate(images, 1):
                completed = subprocess.run(
                    ["tesseract", str(image), "stdout", "-l", "kor+eng", "--psm", "3"],
                    check=False, capture_output=True, timeout=120,
                )
                page_text = completed.stdout.decode("utf-8", errors="replace").strip()
                if page_text:
                    chunks.append(f"[OCR 페이지 {number}]\n{page_text}")
                if not _full_text.get() and sum(map(len, chunks)) >= MAX_EXTRACTED_CHARS:
                    break
        except (subprocess.SubprocessError, OSError) as exc:
            raise ParseError(f"스캔 PDF OCR에 실패했습니다: {type(exc).__name__}") from exc
    if not chunks:
        raise ParseError("OCR에서도 읽을 수 있는 텍스트를 찾지 못했습니다.")
    return "\n\n".join(chunks)

def _docx(path: Path) -> str:
    with _safe_zip(path):
        doc = Document(str(path))
        chunks = [p.text for p in doc.paragraphs]
        for table in doc.tables:
            chunks.extend(" | ".join(cell.text for cell in row.cells) for row in table.rows)
        return "\n".join(chunks)

def _xlsx(path: Path) -> str:
    with _safe_zip(path):
        book = load_workbook(path, read_only=True, data_only=True, keep_links=False)
        chunks: list[str] = []
        try:
            for sheet in book.worksheets:
                chunks.append(f"[시트: {sheet.title}]")
                for row in sheet.iter_rows(values_only=True):
                    values = [str(value) for value in row if value not in (None, "")]
                    if values:
                        chunks.append(" | ".join(values))
                    if not _full_text.get() and sum(map(len, chunks)) >= MAX_EXTRACTED_CHARS:
                        return "\n".join(chunks)
        finally:
            book.close()
        return "\n".join(chunks)

def _pptx(path: Path) -> str:
    with _safe_zip(path):
        deck = Presentation(str(path))
        chunks: list[str] = []
        for number, slide in enumerate(deck.slides, 1):
            chunks.append(f"[슬라이드 {number}]")
            for shape in slide.shapes:
                if hasattr(shape, "text") and shape.text:
                    chunks.append(shape.text)
        return "\n".join(chunks)

def _hwpx(path: Path) -> str:
    chunks: list[str] = []
    with _safe_zip(path) as archive:
        names = sorted(name for name in archive.namelist() if name.lower().endswith(".xml"))
        for name in names:
            root = ElementTree.fromstring(archive.read(name))
            for node in root.iter():
                if node.text and node.text.strip():
                    chunks.append(node.text.strip())
    return "\n".join(chunks)

def _hwp(path: Path) -> str:
    if not olefile.isOleFile(str(path)):
        raise ParseError("유효한 HWP OLE 문서가 아닙니다.")
    chunks: list[str] = []
    with olefile.OleFileIO(str(path)) as ole:
        header = ole.openstream("FileHeader").read()
        compressed = bool(header[36] & 1) if len(header) > 36 else False
        sections = sorted(
            (entry for entry in ole.listdir() if len(entry) == 2 and entry[0] == "BodyText" and entry[1].startswith("Section")),
            key=lambda entry: int(re.sub(r"\D", "", entry[1]) or 0),
        )
        for entry in sections:
            data = ole.openstream(entry).read()
            if compressed:
                data = zlib.decompress(data, -15)
            offset = 0
            while offset + 4 <= len(data):
                header_value = struct.unpack_from("<I", data, offset)[0]
                offset += 4
                tag_id = header_value & 0x3FF
                size = (header_value >> 20) & 0xFFF
                if size == 0xFFF:
                    if offset + 4 > len(data):
                        break
                    size = struct.unpack_from("<I", data, offset)[0]
                    offset += 4
                payload = data[offset : offset + size]
                offset += size
                if tag_id == 67:
                    chunks.append(payload.decode("utf-16le", errors="ignore"))
    return "\n".join(chunks)

def _plain(path: Path, extension: str) -> str:
    raw = path.read_bytes()
    text = raw.decode("utf-8-sig", errors="replace")
    if extension == ".csv":
        return "\n".join(" | ".join(row) for row in csv.reader(io.StringIO(text)))
    return text

def _zip_documents(path: Path) -> str:
    supported = {".pdf", ".docx", ".xlsx", ".pptx", ".txt", ".md", ".csv", ".hwp", ".hwpx"}
    chunks: list[str] = []
    with _safe_zip(path) as archive, tempfile.TemporaryDirectory(prefix="kodame-zip-") as temp_dir:
        temp_root = Path(temp_dir)
        candidates = [item for item in archive.infolist() if not item.is_dir() and Path(item.filename).suffix.lower() in supported]
        if not candidates:
            raise ParseError("ZIP 안에서 분석 가능한 문서를 찾지 못했습니다.")
        for index, item in enumerate(candidates, 1):
            extension = Path(item.filename).suffix.lower()
            extracted = temp_root / f"{index}{extension}"
            extracted.write_bytes(archive.read(item))
            try:
                inner_text, method = parse_document(extracted, extension)
                chunks.append(f"[ZIP 문서: {Path(item.filename).name} | {method}]\n{inner_text}")
            except ParseError as exc:
                if _full_text.get():
                    raise ParseError(f"ZIP 내부 문서 전체 분석 실패: {Path(item.filename).name}: {exc}") from exc
                chunks.append(f"[ZIP 문서 파싱 제외: {Path(item.filename).name} | {exc}]")
            if not _full_text.get() and sum(map(len, chunks)) >= MAX_EXTRACTED_CHARS:
                break
    return "\n\n".join(chunks)

def parse_document(path: Path, extension: str, *, full_text: bool = False) -> tuple[str, str]:
    token = _full_text.set(full_text or _full_text.get())
    try:
        return _parse_document(path, extension)
    finally:
        _full_text.reset(token)


def _parse_document(path: Path, extension: str) -> tuple[str, str]:
    parsers = {
        ".pdf": (_pdf, "pypdf-or-kor-eng-ocr"), ".docx": (_docx, "python-docx"),
        ".xlsx": (_xlsx, "openpyxl-readonly"), ".pptx": (_pptx, "python-pptx"),
        ".hwpx": (_hwpx, "hwpx-defusedxml"), ".hwp": (_hwp, "hwp-ole-records"),
        ".txt": (lambda p: _plain(p, extension), "utf8-text"),
        ".md": (lambda p: _plain(p, extension), "utf8-text"),
        ".csv": (lambda p: _plain(p, extension), "csv"),
        ".zip": (_zip_documents, "safe-zip-recursive"),
    }
    if extension not in parsers:
        raise ParseError(f"지원하지 않는 파일 형식입니다: {extension}")
    parser, method = parsers[extension]
    try:
        text = _trim(parser(path))
    except ParseError:
        raise
    except Exception as exc:
        raise ParseError(f"문서 파싱에 실패했습니다: {type(exc).__name__}") from exc
    if not text:
        raise ParseError("문서에서 읽을 수 있는 텍스트를 찾지 못했습니다.")
    return text, method
