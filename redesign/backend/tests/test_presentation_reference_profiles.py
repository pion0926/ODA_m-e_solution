import unittest
from copy import deepcopy
from io import BytesIO

from PIL import Image
from pptx import Presentation
from pydantic import ValidationError

from kodame_intake.presentation_profiles import get_profile, PresentationRequest
from kodame_intake.presentation_reference_prompt import validate_batch, batch_prompt
from kodame_intake.presentation_reference_renderer import build_reference_deck
from kodame_intake.presentation_quality import render_and_validate_presentation


def fixture(count):
    profile = get_profile(count)
    ids = {k for p in profile['pages'] for k in p['source_sections']}
    source = {'summary': {'project': {'title': '우즈베키스탄 응급의료 교육사업', 'period': '2022~2029'}},
              'report_sections': [{'part_id': k, 'content': '현재 사업에서 제도적 운영 기반을 확인함'} for k in ids],
              'structured_evidence': {}, 'evidence_catalog': []}
    slides = []
    for p in profile['pages']:
        slides.append({'slide_number': p['slide_number'], 'title': p['title'],
                       'rows': [['확인된 근거' for c in p['columns']] for _ in range(5)] if p['columns'] else [],
                       'blocks': [] if p['columns'] else [{'heading': '제도적 운영 기반', 'text': '등록된 자료에서 교육과정의 운영 기반을 확인함. 추가 실적은 후속 확인이 필요함.'}],
                       'photo_id': '', 'caption': '', 'speaker_notes': '현재 자료의 확인 범위와 한계를 구분함.',
                       'source_sections': p['source_sections'][:1]})
    return profile, source, slides


class ReferenceProfilesTest(unittest.TestCase):
    def test_only_fifteen_and_thirty_are_accepted(self):
        for count in (15, 30):
            self.assertEqual(PresentationRequest(slide_count=count).slide_count, count)
        for count in (10, 12, 50, 0, -1):
            with self.assertRaises(ValidationError):
                PresentationRequest(slide_count=count)

    def test_reference_mapping_covers_both_complete_sources(self):
        for count, reference_count in ((15, 15), (30, 50)):
            p = get_profile(count)
            self.assertEqual(len(p['pages']), count)
            self.assertEqual(set(range(1, reference_count+1)), {i for s in p['pages'] for i in s['source_pages']})

    def test_rejects_missing_page_wrong_table_and_unknown_photo(self):
        p, source, slides = fixture(15)
        for mutation in ('missing', 'columns', 'photo'):
            broken = deepcopy(slides)
            if mutation == 'missing': broken.pop()
            if mutation == 'columns': broken[2]['rows'][0].append('추가 열')
            if mutation == 'photo': broken[11]['photo_id'] = 'foreign-project'
            with self.assertRaises(ValueError):
                validate_batch({'slides': broken}, p['pages'], source, {})

    def test_editable_tables_counts_ratios_notes_and_pdf_render(self):
        for count in (15, 30):
            with self.subTest(count=count):
                p, source, slides = fixture(count)
                validate_batch({'slides': slides}, p['pages'], source, {})
                data = build_reference_deck(p, slides, source, {})
                deck = Presentation(BytesIO(data))
                self.assertEqual(len(deck.slides), count)
                self.assertAlmostEqual(deck.slide_width / deck.slide_height, p['width_pt'] / 540)
                self.assertTrue(any(s.has_table for sl in deck.slides for s in sl.shapes))
                for sl in deck.slides:
                    self.assertIn('[Sources]', sl.notes_slide.notes_text_frame.text)
                plan = {'slide_count': count, 'slides': [dict(s, layout=x['layout']) for s, x in zip(slides,p['pages'])]}
                self.assertEqual(render_and_validate_presentation(data, plan)['rendered_slide_count'], count)

    def test_photograph_is_contained_with_original_aspect_ratio(self):
        p, source, slides = fixture(30)
        b = BytesIO()
        Image.new('RGB', (900,600), 'gray').save(b, 'JPEG')
        photos = {'p1': {'data': b.getvalue(), 'width': 900, 'height': 600, 'file_name': '현장.pdf', 'page': 2}}
        slides[13]['photo_id'] = 'p1'
        slides[13]['caption'] = '현재 프로젝트 등록 현장 자료'
        deck = Presentation(BytesIO(build_reference_deck(p, slides, source, photos)))
        pictures = [s for s in deck.slides[13].shapes if s.shape_type == 13]
        self.assertEqual(len(pictures), 1)
        self.assertAlmostEqual(pictures[0].width / pictures[0].height, 1.5, places=4)
        self.assertIn('현장.pdf', deck.slides[13].notes_slide.notes_text_frame.text)

    def test_six_rows_at_the_prompt_budget_fit_every_table(self):
        for count in (15, 30):
            profile, source, slides = fixture(count)
            for p, s in zip(profile['pages'], slides):
                if p['columns']:
                    s['rows'] = [['가'*n for n in p['cell_char_limits']] for _ in range(6)]
            build_reference_deck(profile, slides, source, {})


if __name__ == '__main__':
    unittest.main()
