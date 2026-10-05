import unittest
from io import BytesIO
from zipfile import ZipFile
from openpyxl import load_workbook
from kodame_intake.submission_exports import grade_workbook, feedback_workbook, QUESTION_ROWS
from kodame_intake.dac_rules import mean_score
from kodame_intake.submission_presentations import feedback_deck, lesson_deck


def criteria(score=2):
    return [{'id':cid,'score':score,'question_assessments':[
        {'question_id':f'{cid}-q{i}','score':score,'finding':'문서에 확인된 사실에 따른 판단 근거','limitations':['측정 기간 확인 필요']}
        for i in range(1,len(rows)+1)]} for cid,(rows,_) in QUESTION_ROWS.items()]


FEEDBACK='''1. 우선 성과자료 정비
구분: 자료관리
제언: 성과지표별 원자료를 분기별로 수집하고 검토한다.
이해관계자: 사업수행팀
선정 사유: 동일 지표의 기간별 근거자료 누락이 확인되어 후속 점검이 필요함.
우선순위: 상
완료기한: 후속 계획 확정 시
점검주기: 분기
후속 확인자료: 원자료 목록과 점검 기록
'''
LESSONS='''❍ (성과자료의 측정기간 일치)
성과지표의 정의와 측정기간을 사전에 맞추어야 사업 종료 시 목표와 실적을 같은 기준으로 비교할 수 있음. 자료의 비교가능성을 확보한 유사 사업에 적용할 수 있음.
- 체크리스트 질문: 성과지표의 측정기간과 정의를 사전에 합의했는가?
'''


class SubmissionTests(unittest.TestCase):
    def test_narrative_feedback_keeps_each_full_record(self):
        from backend.oda_me.hwpx.patchers import parse_feedback_items
        body='ㅇ 과제 개요\n설명\nㅇ 첫 번째 과제\n'+FEEDBACK.replace('1. 우선 성과자료 정비','').replace('\n','\n- ')
        body+='\n\nㅇ 두 번째 과제\n'+FEEDBACK.replace('1. 우선 성과자료 정비','').replace('성과지표별 원자료를 분기별로 수집하고 검토한다.','두 번째 조치를 실시한다.').replace('\n','\n- ')
        rows=parse_feedback_items(body)
        self.assertEqual(len(rows),2)
        self.assertEqual(rows[0]['owner'],'사업수행팀')
        self.assertEqual(rows[0]['priority'],'상')
        self.assertEqual(rows[1]['task'],'두 번째 조치를 실시한다.')
    def test_grade_boundary_gap_fixed_and_cached_values_match(self):
        data=grade_workbook('=SUM(1,2) 테스트 사업',criteria(1.7))
        cached=load_workbook(BytesIO(data),data_only=True)
        formulas=load_workbook(BytesIO(data),data_only=False)
        self.assertEqual(cached.worksheets[0]['D58'].value,8.5)
        self.assertEqual(cached.worksheets[0]['D59'].value,'미흡')
        self.assertEqual(formulas.worksheets[0]['C4'].data_type,'s')
        self.assertEqual(cached.worksheets[0]['E9'].value,'/4점')
        self.assertIn('판정 근거',cached.sheetnames)
        self.assertIn('판단 근거',cached['판정 근거']['C4'].value)
        self.assertEqual(cached.worksheets[1]['D63'].value,'미입력')
        self.assertNotIn('8>D58',formulas.worksheets[0]['D59'].value)

    def test_unscored_is_not_zero_and_partial_mean_is_rejected(self):
        data=grade_workbook('사업',criteria(None))
        sheet=load_workbook(BytesIO(data),data_only=True).worksheets[0]
        self.assertEqual(sheet['D58'].value,'판정보류')
        self.assertEqual(sheet['D59'].value,'판정보류')
        inconsistent=criteria();inconsistent[0]['question_assessments'][0]['score']=None
        with self.assertRaises(ValueError):grade_workbook('사업',inconsistent)

    def test_feedback_retains_headers_removes_examples_macros_and_blank_print_pages(self):
        data=feedback_workbook('현재 사업',FEEDBACK)
        with ZipFile(BytesIO(data)) as z:
            self.assertFalse(any('vbaProject' in n for n in z.namelist()))
            self.assertNotIn(b'macroEnabled',z.read('[Content_Types].xml'))
        sheet=load_workbook(BytesIO(data),data_only=True).worksheets[0]
        self.assertEqual(sheet['B4'].value,'현재 사업')
        self.assertEqual(sheet['K4'].value,'면담 실시 여부 확인 필요')
        self.assertNotIn('C사업',' '.join(str(c.value) for row in sheet for c in row))
        self.assertIn('$L$4',str(sheet.print_area))

    def test_feedback_ppt_has_all_columns_and_no_false_confirmation(self):
        from pptx import Presentation
        prs=Presentation(BytesIO(feedback_deck('현재 사업',FEEDBACK)))
        text='\n'.join(s.text for slide in prs.slides for s in slide.shapes if s.has_text_frame)
        self.assertIn('면담 실시 여부 확인 필요',text)
        self.assertIn('상충관계',text)
        self.assertIn('점검주기',text)

    def test_lesson_template_examples_do_not_survive(self):
        data=lesson_deck('현재 사업',LESSONS,['사업 성과자료'])
        with ZipFile(BytesIO(data)) as z:
            text='\n'.join(z.read(n).decode() for n in z.namelist() if n.startswith('ppt/slides/slide') and n.endswith('.xml'))
        for stale in ['2023년','2022.00.00','파라과이','파키스탄','XXXX','설명 설명','수행기관 파트너십 강화']:
            self.assertNotIn(stale,text)
