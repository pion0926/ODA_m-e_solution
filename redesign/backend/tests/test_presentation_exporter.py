from __future__ import annotations

import unittest
from io import BytesIO

from pptx import Presentation

from kodame_intake.presentation_exporter import (
    LAYOUT_SEQUENCE,
    SLIDE_COUNT,
    _ordered_pdm_steps,
    _presentation_prompt,
    build_presentation_bytes,
    normalize_presentation_plan,
    validate_presentation_bytes,
)
from kodame_intake.presentation_quality import render_and_validate_presentation, score_presentation_quality
from kodame_intake.presentation_reference import build_reference_context, select_reference_variant
from kodame_intake.theory_visual import (
    build_theory_png_bytes,
    build_theory_pptx_bytes,
    normalize_theory_visual_plan,
)


def raw_plan(*, repeated_layout: bool = False) -> dict:
    slides = []
    for index in range(SLIDE_COUNT):
        slides.append({
            "slide_number": index + 1,
            "role": "draft",
            "layout": "statement" if repeated_layout else LAYOUT_SEQUENCE[index],
            "eyebrow": f"SECTION {index + 1}",
            "title": f"슬라이드 {index + 1}의 핵심 판단",
            "headline": f"보고서 근거를 압축한 메시지 {index + 1}",
            "body": "등록된 사업자료와 평가결과를 기준으로 핵심 내용을 설명한다.",
            "bullets": ["확인된 근거", "핵심 판단", "후속 시사점"],
            "secondary_bullets": ["강점", "제약", "확인 과제"],
            "metrics": [
                {"value": "15.5", "label": "종합점수", "context": "20점 만점"},
                {"value": "5개", "label": "DAC 평가기준", "context": "질문별 4점 척도"},
            ],
            "metric_value": "3.2",
            "metric_label": "4점 척도",
            "quote": "제도화와 운영역량을 함께 관리해야 성과가 지속된다.",
            "steps": ["활동", "산출물", "성과", "영향"],
            "accent": "teal",
            "visual_direction": "주장과 근거의 위계를 살리고 충분한 여백을 둔다.",
            "speaker_notes": "확인된 근거와 제한사항을 구분하여 설명한다.",
            "source_sections": ["summary-ko", "conclusion"],
            "source_documents": ["사업계획서.pdf"],
        })
    return {
        "deck_title": "종료평가 결과 브리핑",
        "subtitle": "우즈베키스탄 응급구조학과 사업",
        "design": {"palette": "ocean", "mood": "차분한 편집 디자인", "design_rationale": "슬라이드 역할마다 실루엣을 달리한다."},
        "slides": slides,
    }


CONTEXT = {
    "project": {
        "title": "우즈베키스탄 안디잔주 의과대학 응급구조학과 신설 및 교육프로그램 구축",
        "country": "우즈베키스탄",
        "period": "2022-2029",
        "budget": "145만불",
    },
    "overall": {"score": 15.5, "maxScore": 20, "koicaGrade": "B"},
    "criteria": [
        {"name": "적절성", "currentScore4": 3.5},
        {"name": "일관성", "currentScore4": 3.0},
        {"name": "효과성", "currentScore4": 3.2},
        {"name": "효율성", "currentScore4": 2.8},
        {"name": "지속가능성", "currentScore4": 3.0},
    ],
}

REFERENCE_PROFILE = {
    "profile_version": "test.1",
    "reference_folder": {
        "url": "https://drive.google.com/drive/folders/13UG78hLEcJMfWNQo0D5Nc5vMZbtvZahr",
        "content_policy": "reference_format_only",
    },
    "content_firewall": {
        "allowed": ["양식과 흐름"],
        "forbidden": ["샘플 내용 복사"],
    },
    "narrative_patterns": ["판단에서 근거와 조치로 전개"],
    "visual_patterns": ["일관된 그리드와 다양한 실루엣"],
    "variant_families": [
        {"id": "editorial_axis", "palette": "ocean"},
        {"id": "field_brief", "palette": "forest"},
        {"id": "evidence_signal", "palette": "ink"},
        {"id": "cobalt_report", "palette": "cobalt"},
    ],
    "quality_target": 90,
    "max_generation_attempts": 3,
}


class PresentationExporterTests(unittest.TestCase):
    def test_prompt_requires_twelve_varied_audience_facing_slides_and_evidence(self) -> None:
        source = {
            "summary": {"project": CONTEXT["project"]},
            "report_sections": [{"part_id": "summary-ko", "content": "저장된 보고서 본문"}],
            "structured_evidence": {"pdm": {}, "evaluation_run": {}},
            "evidence_catalog": [{"file_name": "사업계획서.pdf", "summary": "사업 근거"}],
        }
        reference = build_reference_context(source, "prompt-test", profile=REFERENCE_PROFILE)
        prompt = _presentation_prompt(source, reference)
        self.assertIn("정확히 12장", prompt)
        self.assertIn("최소 9개 이상의 레이아웃", prompt)
        self.assertIn("HWPX 조판 전 원문", prompt)
        self.assertIn("evidence_catalog", prompt)
        self.assertIn("16pt 본문", prompt)
        self.assertIn("few_shot_format_only", prompt)
        self.assertIn("투입·활동:", prompt)
        self.assertIn("13UG78hLEcJMfWNQo0D5Nc5vMZbtvZahr", prompt)
        self.assertIn("format_and_flow_only", prompt)
        self.assertIn("샘플 내용 복사", prompt)
        self.assertIn(reference["selected_variant"]["id"], prompt)

    def test_reference_variant_is_reproducible_and_export_key_controls_diversity(self) -> None:
        source = {"summary": {"project": CONTEXT["project"]}}
        first = select_reference_variant(REFERENCE_PROFILE, source, "export-a")
        repeated = select_reference_variant(REFERENCE_PROFILE, source, "export-a")
        candidates = {
            select_reference_variant(REFERENCE_PROFILE, source, f"export-{index}")["id"]
            for index in range(12)
        }
        self.assertEqual(first, repeated)
        self.assertGreaterEqual(len(candidates), 3)

    def test_pdm_flow_repairs_a_missing_input_activity_stage(self) -> None:
        stages = _ordered_pdm_steps({
            "steps": [
                "산출: 표준 교육과정과 실습환경",
                "성과: 현지 교원 단독 운영",
                "영향: 병원 전 응급의료 역량 변화",
            ]
        })
        self.assertEqual(len(stages), 4)
        self.assertTrue(stages[0].startswith("투입·활동:"))
        self.assertTrue(stages[1].startswith("산출"))
        self.assertTrue(stages[2].startswith("성과:"))
        self.assertTrue(stages[3].startswith("영향:"))

    def test_repetitive_claude_layouts_are_repaired(self) -> None:
        raw = raw_plan(repeated_layout=True)
        raw["deck_title"] = CONTEXT["project"]["title"]
        plan = normalize_presentation_plan(raw, CONTEXT["project"])
        layouts = [slide["layout"] for slide in plan["slides"]]
        self.assertEqual(tuple(layouts), LAYOUT_SEQUENCE)
        self.assertEqual(len(set(layouts)), SLIDE_COUNT)
        self.assertEqual(plan["deck_title"], "종료평가 결과 브리핑")
        self.assertEqual([slide["role"] for slide in plan["slides"]][0], "cover")
        self.assertEqual([slide["role"] for slide in plan["slides"]][-1], "closing")

    def test_pptx_builder_exports_twelve_editable_slides_with_sources_notes(self) -> None:
        source = {"summary": {"project": CONTEXT["project"]}}
        reference = build_reference_context(source, "quality-test", profile=REFERENCE_PROFILE)
        plan = normalize_presentation_plan(raw_plan(), CONTEXT["project"], reference_context=reference)
        data = build_presentation_bytes(plan, CONTEXT)
        validation = validate_presentation_bytes(data, plan)
        self.assertEqual(validation["slide_count"], SLIDE_COUNT)
        self.assertEqual(validation["unique_layouts"], SLIDE_COUNT)
        presentation = Presentation(BytesIO(data))
        self.assertEqual(len(presentation.slides), SLIDE_COUNT)
        visible = " ".join(
            shape.text for shape in presentation.slides[0].shapes if hasattr(shape, "text")
        )
        self.assertIn("종료평가 결과", visible)
        notes = presentation.slides[0].notes_slide.notes_text_frame.text
        self.assertIn("[Sources]", notes)
        self.assertIn("평가보고서", notes)
        quality = score_presentation_quality(
            data,
            plan,
            validation,
            {"rendered_slide_count": SLIDE_COUNT, "rendered_titles_verified": SLIDE_COUNT},
            reference,
        )
        self.assertGreaterEqual(quality["quality_score"], 90)
        self.assertTrue(quality["passed"])

    def test_pptx_renders_all_slides_and_titles_without_loss(self) -> None:
        plan = normalize_presentation_plan(raw_plan(), CONTEXT["project"])
        data = build_presentation_bytes(plan, CONTEXT)
        validation = render_and_validate_presentation(data, plan)
        self.assertEqual(validation["rendered_slide_count"], SLIDE_COUNT)
        self.assertEqual(validation["rendered_titles_verified"], SLIDE_COUNT)

    def test_theory_visual_builds_one_slide_pptx_and_matching_wide_png(self) -> None:
        plan = normalize_theory_visual_plan({
            "title": "변화이론 분석",
            "subtitle": "교육체계 구축과 현지 운영역량 강화가 응급의료 서비스 개선으로 연결된다.",
            "challenges": ["병원 전 단계 대응 취약", "전문인력 부족", "표준 교육체계 미비"],
            "activity_groups": [
                {"label": "학과 기반", "items": ["학과 승인", "교육과정 개발"]},
                {"label": "현장 역량", "items": ["교원 양성", "지역사회 교육"]},
            ],
            "pre_output_factors": [
                {"kind": "working", "text": "관계부처 승인과 대학 협력"},
                {"kind": "nonworking", "text": "초기 대면 수요조사 제약"},
            ],
            "outputs": ["정규 학과 운영", "현지 교원 양성", "표준 교재와 실습환경"],
            "post_output_factors": [
                {"kind": "working", "text": "현지 교원 단독 강의 이행"},
                {"kind": "nonworking", "text": "졸업생 성과 산정 전"},
            ],
            "outcomes": ["전문인력 배출 기반", "지역사회 대응역량 강화", "응급의료 서비스 개선"],
        })
        pptx = build_theory_pptx_bytes(plan)
        png = build_theory_png_bytes(plan)
        presentation = Presentation(BytesIO(pptx))
        self.assertEqual(len(presentation.slides), 1)
        self.assertTrue(any("변화이론" in shape.text for shape in presentation.slides[0].shapes if hasattr(shape, "text")))
        visible = " ".join(shape.text for shape in presentation.slides[0].shapes if hasattr(shape, "text"))
        for label in ("당면과제", "주요 활동", "작동·비작동요인", "산출물", "중장기성과"):
            self.assertIn(label, visible)
        self.assertTrue(png.startswith(b"\x89PNG"))


if __name__ == "__main__":
    unittest.main()
