"""Compile the designated v1.4 workbook into a versioned, reviewable rule set."""
import hashlib
import json
import re
from pathlib import Path
from openpyxl import load_workbook

root = Path(__file__).resolve().parents[1]
path = next((root / 'samples/dac-scoring-20260918').glob('*.xlsm'))
book = load_workbook(path, data_only=False, keep_vba=True)
sheet, detail, settings = book['종료평가용'], book['세부평가 입력'], book['기준 설정']
ids = ['relevance-q1','relevance-q2','coherence-q1','coherence-q2',
       'effectiveness-q1','effectiveness-q2','effectiveness-q3',
       'efficiency-q1','efficiency-q2','sustainability-q1','sustainability-q2']
questions = {}
official_rows = {re.match(r'^(\d\.\d)', str(sheet[f'C{r}'].value)).group(1): r
                 for r in range(1, sheet.max_row + 1)
                 if re.match(r'^\d\.\d', str(sheet[f'C{r}'].value))}
for n, qid in enumerate(ids):
    row = 9 + n * 5
    official_row = official_rows[sheet[f'U{row}'].value[1:]]
    questions[qid] = {
        'source_question_id': sheet[f'U{row}'].value,
        'gate': sheet[f'AC{row}'].value.split('4점 Gate: ')[1].split('\n상한:')[0],
        'cap_condition': settings[f'C{22+n}'].value,
        'red_flag_condition': settings[f'D{22+n}'].value,
        'official_levels': {str(level): sheet[f'G{official_row+level-1}'].value for level in range(1,5)},
        'level_requirements': {str(level): sheet[f'H{official_row+level-1}'].value for level in range(1,5)},
        'checks': [{'id': detail[f'C{6+n*5+i}'].value,
                    'criterion': sheet[f'X{row+i}'].value, 'weight': 20,
                    'min_grade_for_verified': 3 if detail[f'C{6+n*5+i}'].value in {
                        'CQ1_I05','CQ2_I05','EQ1_I01','EQ2_I01','EQ2_I03','EQ2_I04',
                        'EQ3_I04','EQ3_I05','FQ1_I04','FQ1_I05'} else 2,
                    'required_evidence': sheet[f'AA{row+i}'].value,
                    'source_locator': f'종료평가용!X{row+i}:AA{row+i}'} for i in range(5)]}
config = {
    'version': 'odame-ears-1.4.20260918.1',
    'label': '제공된 ODA-EARS v1.4 기반 증빙 평가',
    'notice': '제공된 KOICA 종료평가 양식 및 내부 연구자료를 적용한 진단이며 기관의 공식 확정·서명 평가가 아닙니다.',
    'source': {'file': path.name, 'sha256': hashlib.sha256(path.read_bytes()).hexdigest(),
               'folder': 'https://drive.google.com/drive/folders/1Uw3aqKFflx9jcuzTXezdaUscWNwvuUqJ'},
    'adaptation': '반영용 55개 항목·동일가중·임계값을 사용. 참고 개발명세 5~7장에 따라 성과판정과 증거품질을 분리하며 품질 체크 수로 성과를 만들지 않음. 질문 점수는 정수, 기준 평균은 소수점 한 자리 HALF_UP. 미확인·충돌은 null.',
    'states': {'negative': 0, 'limited': 0.33, 'substantial': 0.67, 'verified': 1, 'unverified': None, 'conflicted': None},
    'merit_thresholds': [40,60,80], 'coverage_min': 0.6, 'coverage_four': 0.85,
    'valid_dq_min': 0.5, 'confidence_min': 50,
    'dq_weights': {'reliability':0.3,'directness':0.25,'traceability':0.2,'triangulation':0.15,'recency':0.1},
    'questions': questions}
target = root / 'config/dac_evidence_rubric.json'
target.write_text(json.dumps(config, ensure_ascii=False, indent=2)+'\n', encoding='utf-8')
print(f'{len(questions)} questions / {sum(len(q["checks"]) for q in questions.values())} indicators -> {target}')
