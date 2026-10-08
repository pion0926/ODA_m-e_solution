"""Review the app renderer against stored plans without changing service exports."""
import json
from io import BytesIO
from pathlib import Path
from pptx import Presentation
from PIL import Image
from kodame_intake.presentation_profiles import get_profile
from kodame_intake.presentation_reference_renderer import build_reference_deck

root = Path(__file__).resolve().parents[1]
out = root / 'output/review-20260916'
out.mkdir(parents=True, exist_ok=True)
for count, export_id in [(15,'7bfdd911-cc8b-4d2f-9a86-0666e408fad3'),(30,'574cf3bd-40e8-4cc1-9dc9-0f369f6fcdf3')]:
    base = root/'data-redesign/presentation_exports'/export_id
    plan = json.loads(base.with_suffix('.json').read_text(encoding='utf-8'))
    deck = Presentation(base.with_suffix('.pptx'))
    photos = {}
    for item, slide in zip(plan['slides'],deck.slides):
        key = item.get('photo_id')
        if key:
            pic = next(s for s in slide.shapes if s.shape_type==13)
            im = Image.open(BytesIO(pic.image.blob))
            label = slide.notes_slide.notes_text_frame.text.rsplit('\n사진: ',1)[-1]
            name,page = label.rsplit(', p. ',1)
            photos[key] = {'data':pic.image.blob,'width':im.width,'height':im.height,'file_name':name,'page':int(page)}
    source = {'summary':plan['source_summary']}
    data = build_reference_deck(get_profile(count),plan['slides'],source,photos)
    target = out/f'ODAME-layout-review-{count}.pptx'
    target.write_bytes(data)
    print(target)
