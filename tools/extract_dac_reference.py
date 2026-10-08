"""Read the supplied scoring sources without executing workbook macros."""
from pathlib import Path
import hashlib
import json
from docx import Document
from openpyxl import load_workbook
from pypdf import PdfReader

root = Path(__file__).resolve().parents[1] / 'samples/dac-scoring-20260918'
out = root / 'extracted'
out.mkdir(exist_ok=True)
manifest = []
for path in sorted(root.iterdir()):
    if not path.is_file() or path.suffix not in {'.docx', '.xlsm', '.pdf'}:
        continue
    manifest.append({'name': path.name, 'bytes': path.stat().st_size,
                     'sha256': hashlib.sha256(path.read_bytes()).hexdigest()})
    if path.suffix == '.docx':
        doc = Document(path)
        lines = []
        for element in doc.element.body:
            if element.tag.endswith('}p'):
                from docx.text.paragraph import Paragraph
                lines.append(Paragraph(element, doc).text)
            elif element.tag.endswith('}tbl'):
                from docx.table import Table
                for row in Table(element, doc).rows:
                    lines.append(' | '.join(c.text for c in row.cells))
        (out / (path.stem + '.txt')).write_text('\n'.join(lines), encoding='utf-8')
    elif path.suffix == '.xlsm':
        book = load_workbook(path, keep_vba=True, data_only=False)
        values = load_workbook(path, keep_vba=True, data_only=True)
        sheets = {}
        for sheet in book:
            rows = []
            for row in sheet:
                cells = [{'cell': c.coordinate, 'value': c.value,
                          'cached': values[sheet.title][c.coordinate].value}
                         for c in row if c.value is not None]
                if cells:
                    rows.append(cells)
            sheets[sheet.title] = {'rows': rows, 'merged': list(map(str, sheet.merged_cells.ranges))}
        (out / 'workbook.json').write_text(json.dumps(sheets, ensure_ascii=False, indent=2, default=str), encoding='utf-8')
    elif path.suffix == '.pdf':
        reader = PdfReader(path)
        (out / 'rating-table.txt').write_text('\n'.join(f'PAGE {i+1}\n{p.extract_text()}' for i,p in enumerate(reader.pages)), encoding='utf-8')
(root / 'manifest.json').write_text(json.dumps(manifest, ensure_ascii=False, indent=2), encoding='utf-8')
print(json.dumps(manifest, ensure_ascii=False))
