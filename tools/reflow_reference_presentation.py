"""Re-render an explicitly named completed export without another model call.

Run in the API container. Keeps original embedded photographs, archives the
previous PPTX, and only replaces the export after real-render validation.
"""
import json
import sys
from io import BytesIO
from pathlib import Path
from uuid import UUID

from PIL import Image
from pptx import Presentation
from psycopg.types.json import Jsonb
from kodame_intake.db import connection, tenant_context, pool
from kodame_intake.settings import DATA_DIR
from kodame_intake.presentation_source import collect_presentation_source
from kodame_intake.presentation_reference_renderer import build_reference_deck
from kodame_intake.presentation_quality import render_and_validate_presentation


def reflow(export_id):
    export_id = str(UUID(export_id))
    with tenant_context(None, system=True), connection() as conn:
        row = conn.execute('SELECT project_id,status,validation FROM presentation_exports WHERE id=%s', (export_id,)).fetchone()
    if not row or row['status'] != 'completed':
        raise RuntimeError('Only an explicitly named completed export can be reflowed.')
    directory = DATA_DIR / 'presentation_exports'
    path = directory / f'{export_id}.pptx'
    plan_path = directory / f'{export_id}.json'
    plan = json.loads(plan_path.read_text(encoding='utf-8'))
    original = path.read_bytes()
    deck = Presentation(BytesIO(original))
    photos = {}
    for item, slide in zip(plan['slides'], deck.slides):
        key = item.get('photo_id')
        if not key:
            continue
        pictures = [s for s in slide.shapes if s.shape_type == 13]
        if len(pictures) != 1:
            raise RuntimeError(f'Cannot safely identify original photograph on slide {item["slide_number"]}.')
        data = pictures[0].image.blob
        im = Image.open(BytesIO(data))
        note = slide.notes_slide.notes_text_frame.text
        label = note.rsplit('\n사진: ', 1)[-1]
        name, page = label.rsplit(', p. ', 1)
        photos[key] = {'data': data, 'width': im.width, 'height': im.height, 'file_name': name, 'page': int(page)}
    with tenant_context(row['project_id'], system=row['project_id'] is None):
        summary = plan.get('source_summary')
        if summary is None:
            _, source, _ = collect_presentation_source()
            summary = source['summary']
        source = {'summary': summary}
        data = build_reference_deck(plan['profile'], plan['slides'], source, photos)
        check = {'slide_count': plan['profile']['slide_count'], 'slides': [dict(s, layout=p['layout']) for s,p in zip(plan['slides'], plan['profile']['pages'])]}
        validation = render_and_validate_presentation(data, check)
        validation = {**(row['validation'] or {}), **validation, 'layout_review': 'duplicate heading removed; existing content and photos preserved'}
        backup = directory / f'{export_id}.before-layout-review.pptx'
        if not backup.exists():
            backup.write_bytes(original)
        path.write_bytes(data)
        plan['source_summary'] = summary
        plan_path.write_text(json.dumps(plan, ensure_ascii=False, indent=2), encoding='utf-8')
        with connection() as conn, conn.transaction():
            conn.execute('UPDATE presentation_exports SET validation=%s,updated_at=now() WHERE id=%s', (Jsonb(validation), export_id))
    print(json.dumps({'id': export_id, 'validation': validation}, ensure_ascii=False))


if __name__ == '__main__':
    pool.open()
    try:
        for export_id in sys.argv[1:]:
            reflow(export_id)
    finally:
        pool.close()
